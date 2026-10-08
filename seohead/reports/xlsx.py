"""Excel report with Summary, Findings, Pages, and Technologies worksheets.

The chart is built with Excel primitives through openpyxl rather than embedded as
a static matplotlib image. It therefore remains editable and tied to worksheet
data, while the generated workbook avoids an unnecessary plotting dependency.
"""

from __future__ import annotations

import pathlib
from typing import Any

_HEAD = {"critical": "C00000", "warning": "BF8F00", "notice": "808080"}
_MAX_SUPPRESSED_FINDINGS = 10000
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_CELL_CHARS = 32_767
MAX_DATA_ROWS = EXCEL_MAX_ROWS - 1
MAX_WORKSHEETS = 10_000


def _cell(value: Any) -> Any:
    from seohead.reports import neutralize_formula

    value = neutralize_formula(value)
    if isinstance(value, str) and len(value) > EXCEL_MAX_CELL_CHARS:
        raise ValueError(
            "XLSX cell exceeds Excel's 32767-character limit; use complete CSV or JSON output"
        )
    return value


def _append(sheet, values) -> None:
    # Validate before openpyxl silently truncates a string during cell creation.
    sheet.append([_cell(value) for value in values])


def _style_header(ws, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    # Use a string coordinate rather than ``ws.cell()``. Accessing a cell here
    # materializes a row, so the next ``append`` would skip one and leave a blank
    # row immediately below the header.
    ws.freeze_panes = f"A{row + 1}"


def _autofit(ws, limits: dict[int, int] | None = None) -> None:
    """Fit columns to content while capping widths for readable long-URL tables."""
    from openpyxl.utils import get_column_letter

    limits = limits or {}
    for idx, column in enumerate(ws.columns, start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(idx)].width = min(
            max(longest + 2, 10), limits.get(idx, 60)
        )


def _metadata_workbook(document: dict[str, Any]):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font

    from seohead.reports import checks_completed_display, neutralize_formula
    from seohead.reports.client_findings import (
        check_title,
        finding_exclusion_report,
        finding_view_columns,
        finding_view_label,
        finding_view_notice,
    )

    wb = Workbook()
    summary = document.get("summary") or {}
    by_sev = summary.get("findings_by_severity") or {}

    # -- Summary -------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = _cell(f"SEO Audit: {document.get('domain', '')}")
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = _cell(document.get("url", ""))
    ws["A3"] = _cell(f"Generated: {document.get('generated_at', '')}")

    # Failed/partial crawl scope and disabled-check evidence must be visible
    # before the metrics and severity counts below, not appended as a
    # trailing note: a recipient who never scrolls past the numbers must
    # still be unable to mistake a failed or sampled crawl for a clean,
    # site-wide audit, and a deliberately disabled check for one that ran
    # clean (#361).
    scope_rows: list[str] = []
    if summary.get("crawl_valid") is False:
        reason = summary.get("crawl_invalid_reason") or "the crawl produced no usable data"
        scope_rows.append(f"Crawl failed -- no health score. {reason}")
    if summary.get("crawl_partial"):
        finish = summary.get("crawl_finish_reason")
        scope = summary.get("crawl_scope_note")
        bits = [b for b in (f"stopped: {finish}" if finish else None, scope) if b]
        scope_rows.append(
            "Partial crawl -- scope is limited." + (f" {'; '.join(bits)}" if bits else "")
        )
    for item in summary.get("checks_disabled") or []:
        scope_rows.append(f"Disabled check {check_title(item.get('id'))} -- {item.get('reason')}")
    if notice := finding_view_notice(summary):
        scope_rows.append(notice)
    exclusions = finding_exclusion_report(summary, document.get("suppressed_issues"))
    if exclusions is not None:
        count = exclusions["suppressed_total"]
        finding_label = "finding" if count == 1 else "findings"
        occurrences = exclusions["suppressed_occurrences"]
        occurrence_label = "occurrence" if occurrences == 1 else "occurrences"
        rule_count = exclusions["rules_configured"]
        rule_label = "rule" if rule_count == 1 else "rules"
        scope_rows.append(
            f"Finding exclusions: {count} {finding_label} and {occurrences} {occurrence_label} "
            f"suppressed by {rule_count} configured URL {rule_label}; see Finding Exclusions.",
        )

    row = 4
    for text in scope_rows:
        cell = ws.cell(row=row, column=1, value=_cell(text))
        cell.font = Font(bold=True, color="C00000")
        row += 1
    offset = row - 4

    rows = [
        ("Pages checked", summary.get("pages_checked", 0)),
        ("Total findings", summary.get("findings_total", 0)),
        ("Critical findings", by_sev.get("critical", 0)),
        ("Warnings", by_sev.get("warning", 0)),
        ("Notices", by_sev.get("notice", 0)),
        ("Checks completed", checks_completed_display(summary)),
        ("Checks unavailable", len(summary.get("tools_failed") or [])),
    ]
    header_row = 5 + offset
    ws.cell(row=header_row, column=1, value="Metric")
    ws.cell(row=header_row, column=2, value="Value")
    _style_header(ws, header_row)
    for i, (name, value) in enumerate(rows, start=header_row + 1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=_cell(value))

    # The three severity bars expose the issue distribution at a glance.
    sev_start = header_row + 3  # Critical findings is the 3rd metric row
    chart = BarChart()
    chart.title = "Findings by Severity"
    chart.y_axis.title = "Count"
    chart.add_data(
        Reference(ws, min_col=2, min_row=sev_start, max_row=sev_start + 2), titles_from_data=False
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=sev_start, max_row=sev_start + 2))
    chart.legend = None
    chart.height, chart.width = 7, 12
    ws.add_chart(chart, f"D{header_row}")

    failed = summary.get("tools_failed") or []
    if failed:
        start = header_row + 1 + len(rows) + 1
        ws.cell(row=start, column=1, value="Unavailable checks -- evidence is absent from report")
        ws.cell(row=start, column=1).font = Font(bold=True, color="C00000")
        for i, item in enumerate(failed, start=start + 1):
            ws.cell(row=i, column=1, value=_cell(check_title(item.get("tool"))))
            ws.cell(row=i, column=2, value=_cell(item.get("error")))
    note = summary.get("severity_note")
    if note:
        ws.cell(
            row=header_row + 1 + len(rows) + len(failed) + 3, column=1, value=_cell(note)
        ).font = Font(italic=True, size=9, color="808080")
    _autofit(ws, {2: 40})

    if exclusions is not None:
        ws = wb.create_sheet("Finding Exclusions")
        _append(ws, ["Rule", "Pattern", "Checks", "Suppressed findings", "Occurrences", "Reason"])
        _style_header(ws)
        for rule in exclusions["rules"]:
            _append(
                ws,
                [
                    neutralize_formula(rule["id"]),
                    neutralize_formula(rule["pattern"]),
                    neutralize_formula(", ".join(rule["checks"]) or "all checks"),
                    rule["suppressed_findings"],
                    rule["suppressed_occurrences"],
                    neutralize_formula(rule["reason"]),
                ],
            )
        _autofit(ws)

        suppressed = exclusions["issues"]
        if suppressed:
            ws = wb.create_sheet("Suppressed Findings")
            _append(ws, ["Issue ID", "Check", "Severity", "URL", "Occurrences", "Rule", "Reason"])
            _style_header(ws)
            for issue in suppressed[:_MAX_SUPPRESSED_FINDINGS]:
                issue = issue if isinstance(issue, dict) else {}
                marker = issue.get("suppression")
                marker = marker if isinstance(marker, dict) else {}
                _append(
                    ws,
                    [
                        neutralize_formula(issue.get("id", "")),
                        neutralize_formula(check_title(issue.get("check"))),
                        neutralize_formula(issue.get("severity", "")),
                        neutralize_formula(issue.get("target_url", "")),
                        issue.get("occurrences_count", ""),
                        neutralize_formula(marker.get("rule_id", "")),
                        neutralize_formula(marker.get("reason", "")),
                    ],
                )
            if len(suppressed) > _MAX_SUPPRESSED_FINDINGS:
                _append(
                    ws,
                    [
                        f"Showing {_MAX_SUPPRESSED_FINDINGS} of {len(suppressed)}; see source audit JSON for all records"
                    ],
                )
            _autofit(ws)

    # -- Findings ------------------------------------------------------------
    # This sheet is the documented developer handoff for a Screaming Frog
    # audit (docs/scenarios/broken-pages.md): Check/Status/Occurrences/
    # Locations/Fix Hint are the evidence a BROKEN_INTERNAL_LINK finding
    # carries beyond its message, and dropping them here forced the reader
    # back to raw audit.json (#220).
    ws = wb.create_sheet("Findings")
    view_columns = finding_view_columns(summary)
    if view_columns is not None:
        _append(ws, [finding_view_label(column) for column in view_columns])
    else:
        _append(
            ws,
            [
                "Severity",
                "URL",
                "Finding",
                "Observation",
                "Reproduction",
                "Status",
                "Occurrences",
                "Evidence",
                "Locations",
                "Fix Hint",
            ],
        )
    _style_header(ws)
    _autofit(ws, {4: 100, 8: 100, 9: 60} if view_columns is None else {})

    # -- Pages ---------------------------------------------------------------
    ws = wb.create_sheet("Pages")
    # description_length is part of the site-audit page contract (audit.site
    # emits it, csvfile.py already writes it) -- this sheet was the one place
    # it was silently dropped, making the XLSX working file unusable for the
    # meta-description-length scenario it is supposed to cover (#225).
    titles = [
        "URL",
        "Status",
        "Title",
        "Title Length",
        "Description Length",
        "H1",
        "Canonical",
        "Words",
        "Schema Types",
        "Schema Errors",
        "Missing Social Tags",
    ]
    _append(ws, titles)
    _style_header(ws)
    _autofit(ws, {1: 70, 3: 60, 6: 40})  # URL, Title, H1 -- H1 shifted by the new column

    # -- Technologies and infrastructure ------------------------------------
    ws = wb.create_sheet("Technologies")
    _append(ws, ["Category", "Detected Technology", "Evidence"])
    _style_header(ws)
    tech = (document.get("site") or {}).get("tech_detect") or {}
    for item in tech.get("technologies") or []:
        _append(
            ws,
            [
                neutralize_formula(item.get("category", "")),
                neutralize_formula(item.get("name", "")),
                neutralize_formula(item.get("evidence", "")),
            ],
        )
    registration = ((document.get("site") or {}).get("domain_profile") or {}).get(
        "registration"
    ) or {}
    if registration:
        _append(ws, [])
        _append(ws, ["domain", "registrar", neutralize_formula(registration.get("registrar", ""))])
        _append(ws, ["domain", "created", neutralize_formula(registration.get("created", ""))])
        _append(ws, ["domain", "expires", neutralize_formula(registration.get("expires", ""))])
        _append(ws, ["domain", "age in years", registration.get("age_years", "")])
    _autofit(ws, {3: 70})

    # -- Project coverage ----------------------------------------------------
    coverage = summary.get("project_coverage")
    if isinstance(coverage, dict):
        from seohead.reports.project_coverage import value_text

        project = coverage.get("project") or {}
        checklist = coverage.get("status") or {}
        ws = wb.create_sheet("Project Coverage")
        _append(ws, ["Project", neutralize_formula(project.get("site", ""))])
        _append(ws, ["Project UUID", neutralize_formula(project.get("uuid", ""))])
        _append(ws, ["Checklist state", neutralize_formula(checklist.get("state", ""))])
        _append(ws, ["Revision", checklist.get("revision", "")])
        _append(ws, ["Counts", neutralize_formula(value_text(checklist.get("counts")))])
        _append(ws, [])
        _append(
            ws,
            [
                "Item ID",
                "Item",
                "Kind",
                "Execution",
                "Priority",
                "Priority origin",
                "Priority reason",
                "State",
                "Attempt",
                "Complete",
                "Blocked by",
                "Enabled",
                "Stale",
                "Scope",
                "Measurement",
                "Reason",
            ],
        )
        _style_header(ws, 7)
        for item in checklist.get("items") or []:
            if not isinstance(item, dict):
                continue
            _append(
                ws,
                [
                    neutralize_formula(item.get("id", "")),
                    neutralize_formula(item.get("title", "")),
                    neutralize_formula(item.get("kind", "")),
                    neutralize_formula(item.get("execution_kind", "")),
                    neutralize_formula(item.get("priority", "")),
                    neutralize_formula(item.get("priority_origin", "")),
                    neutralize_formula(item.get("priority_reason", "")),
                    neutralize_formula(item.get("state", "")),
                    neutralize_formula(item.get("attempt_status", "")),
                    item.get("complete", ""),
                    neutralize_formula(value_text(item.get("blocked_by"))),
                    item.get("enabled", ""),
                    item.get("stale", ""),
                    neutralize_formula(value_text(item.get("scope"))),
                    neutralize_formula(value_text(item.get("measurement"))),
                    neutralize_formula(item.get("reason", checklist.get("reason", ""))),
                ],
            )
        if ws.max_row > 7:
            ws.auto_filter.ref = f"A7:P{ws.max_row}"
        _autofit(ws, {2: 45, 7: 60, 14: 70, 15: 70, 16: 70})

    from seohead.reports.evidence_summary import rows as evidence_rows

    evidence = evidence_rows(summary)
    if evidence:
        evidence_sheet = wb.create_sheet("Evidence coverage")
        _append(evidence_sheet, ["Kind", "Measurement", "State", "Scope or reason"])
        for row in evidence:
            _append(evidence_sheet, [neutralize_formula(value) for value in row])
        _style_header(evidence_sheet)
        _autofit(evidence_sheet, {2: 80, 4: 100})
        evidence_sheet.auto_filter.ref = evidence_sheet.dimensions

    return wb


def write(document: Any, path: pathlib.Path) -> dict[str, Any]:
    """Use one report layout and stream the two unbounded tables to worksheet XML."""
    import hashlib
    import json
    import os
    import shutil
    import tempfile
    from contextlib import suppress
    from copy import copy
    from itertools import islice

    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from seohead.reports import SEVERITY_TITLES, neutralize_formula
    from seohead.reports.bi import DEFAULT_MAX_OUTPUT_BYTES, MIN_FREE_DISK_BYTES
    from seohead.reports.client_findings import finding_view_columns

    path = pathlib.Path(path)
    index_path = path.with_suffix(path.suffix + ".index.json")
    for candidate in (path, index_path):
        if candidate.is_symlink() or (candidate.exists() and not candidate.is_file()):
            raise OSError("XLSX output and index must be regular files")

    # The small metadata sheets retain the established layout. Only findings and
    # pages grow with the crawl; their rows are replayed to measure widths, then
    # written directly to disk using openpyxl's write-only mode.
    template = _metadata_workbook(
        {
            key: document.get(key)
            for key in ("domain", "url", "generated_at", "summary", "site", "suppressed_issues")
        }
    )
    workbook = Workbook(write_only=True)
    summary = document.get("summary") or {}
    view_columns = finding_view_columns(summary)
    severity_column = (
        (view_columns.index("severity") if "severity" in view_columns else None)
        if view_columns is not None
        else 0
    )

    def finding_rows():
        for finding in document.get("findings") or []:
            if view_columns is not None:
                fields = finding.get("view_fields") or {}
                values = [neutralize_formula(fields.get(column, "")) for column in view_columns]
            else:
                values = [
                    SEVERITY_TITLES.get(finding.get("severity"), finding.get("severity")),
                    neutralize_formula(finding.get("url", "")),
                    neutralize_formula(finding.get("client_title", "Audit finding")),
                    neutralize_formula(finding.get("client_observation", "")),
                    neutralize_formula(finding.get("client_reproduction", "")),
                    finding.get("status_code", ""),
                    finding.get("occurrences_count", ""),
                    neutralize_formula("; ".join(finding.get("client_details") or [])),
                    neutralize_formula("; ".join(finding.get("client_locations") or [])),
                    neutralize_formula(finding.get("fix_hint", "")),
                ]
            yield values, _HEAD.get(finding.get("severity"))

    def page_rows():
        columns = (
            "url",
            "status",
            "title",
            "title_length",
            "description_length",
            "h1",
            "canonical",
            "words",
            "schema_types",
            "schema_errors",
            "social_missing",
        )
        for page in document.get("pages") or []:
            yield [neutralize_formula(page.get(column, "")) for column in columns], None

    ranges = []

    def check_spool() -> None:
        paths = [pathlib.Path(s._writer.out) for s in workbook if s._writer is not None]
        total = sum(p.stat().st_size for p in paths if p.exists())
        if total > DEFAULT_MAX_OUTPUT_BYTES:
            raise ValueError("XLSX spool exceeds its byte bound; use complete CSV or JSON output")
        # openpyxl's XML spools and the final archive can live on different volumes.
        for directory in {path.parent, *(p.parent for p in paths)}:
            if shutil.disk_usage(directory).free < MIN_FREE_DISK_BYTES:
                raise OSError("insufficient free disk for XLSX output; use complete CSV or JSON")

    try:
        check_spool()
        for source in template.worksheets:
            rows = (
                finding_rows
                if source.title == "Findings"
                else (page_rows if source.title == "Pages" else None)
            )
            count = 0
            widths = [len(str(cell.value or "")) for cell in source[1]]
            if rows is not None:
                for values, _colour in rows():
                    count += 1
                    for index, value in enumerate(values):
                        value = _cell(value)
                        if value is not None:
                            widths[index] = max(widths[index], len(str(value)))
            elif source.max_row > EXCEL_MAX_ROWS:
                raise ValueError(
                    "XLSX metadata exceeds Excel's row limit; use complete CSV or JSON"
                )
            parts = max(1, (count + MAX_DATA_ROWS - 1) // MAX_DATA_ROWS)
            if len(workbook.worksheets) + parts > MAX_WORKSHEETS:
                raise ValueError("XLSX worksheet count exceeds its bound; use complete CSV or JSON")
            stream = iter(rows()) if rows is not None else iter(())
            written = 0
            for part in range(parts):
                title = source.title if part == 0 else f"{source.title} {part + 1}"
                sheet = workbook.create_sheet(title)
                for key, dimension in source.column_dimensions.items():
                    sheet.column_dimensions[key] = copy(dimension)
                sheet.freeze_panes = source.freeze_panes
                sheet.auto_filter = copy(source.auto_filter)
                part_rows = min(MAX_DATA_ROWS, count - written)
                if rows is not None:
                    limits = (
                        ({4: 100, 8: 100, 9: 60} if view_columns is None else {})
                        if source.title == "Findings"
                        else {1: 70, 3: 60, 6: 40}
                    )
                    for index, width in enumerate(widths, 1):
                        sheet.column_dimensions[get_column_letter(index)].width = min(
                            max(width + 2, 10), limits.get(index, 60)
                        )
                    if part_rows:
                        sheet.auto_filter.ref = (
                            f"A1:{get_column_letter(len(widths))}{part_rows + 1}"
                        )
                    ranges.append(
                        {
                            "table": source.title,
                            "worksheet": title,
                            "first_record": written + 1 if part_rows else None,
                            "last_record": written + part_rows if part_rows else None,
                            "rows": part_rows,
                            "header_rows": 1,
                        }
                    )
                dimension = f"A1:{get_column_letter(source.max_column)}{source.max_row + part_rows}"
                sheet.calculate_dimension = lambda ref=dimension: ref
                for row in source.iter_rows():
                    cells = []
                    for original in row:
                        cell = WriteOnlyCell(sheet, value=_cell(original.value))
                        cell.font = copy(original.font)
                        cell.fill = copy(original.fill)
                        cell.border = copy(original.border)
                        cell.alignment = copy(original.alignment)
                        cell.number_format = original.number_format
                        cell.protection = copy(original.protection)
                        cells.append(cell)
                    sheet.append(cells)
                for values, colour in islice(stream, part_rows):
                    cells = [WriteOnlyCell(sheet, value=_cell(value)) for value in values]
                    if colour and severity_column is not None and source.title == "Findings":
                        cells[severity_column].font = Font(bold=True, color=colour)
                    sheet.append(cells)
                    written += 1
                    if written % 256 == 1:
                        check_spool()
                for chart in source._charts:
                    sheet.add_chart(chart)
                sheet.close()
            if written != count or next(stream, None) is not None:
                raise ValueError("XLSX row conservation failed; source rows must be replayable")
        check_spool()
        # A failed archive write never touches a previous report. The adjacent
        # hash index is the completion marker; caught publication failures roll back.
        with tempfile.TemporaryDirectory(prefix=".seohead-xlsx-", dir=path.parent) as directory:
            stage = pathlib.Path(directory)
            staged = stage / "report.xlsx"
            workbook.save(staged)
            if staged.stat().st_size > DEFAULT_MAX_OUTPUT_BYTES:
                raise ValueError("XLSX exceeds its byte bound; use complete CSV or JSON output")
            digest = hashlib.sha256()
            with staged.open("rb") as source_file:
                for block in iter(lambda: source_file.read(1024 * 1024), b""):
                    digest.update(block)
            suppressed = document.get("suppressed_issues") or []
            index = {
                "format": "seohead.report-xlsx-index.v1",
                "state": "complete",
                "workbook": path.name,
                "workbook_sha256": digest.hexdigest(),
                "workbook_bytes": staged.stat().st_size,
                "ranges": ranges,
                "finding_view": summary.get("finding_view"),
                "source_population": {
                    "findings": summary.get("findings_total"),
                    "pages": summary.get("pages_checked"),
                },
                "coverage_worksheets": [
                    s.title for s in workbook if s.title not in {r["worksheet"] for r in ranges}
                ],
                "omissions": {
                    "suppressed_findings": max(0, len(suppressed) - _MAX_SUPPRESSED_FINDINGS),
                    "fallback": "Complete CSV scope output or source JSON",
                },
                "limits": {
                    "data_rows_per_sheet": MAX_DATA_ROWS,
                    "cell_characters": EXCEL_MAX_CELL_CHARS,
                    "max_output_bytes": DEFAULT_MAX_OUTPUT_BYTES,
                },
            }
            staged_index = stage / "index.json"
            staged_index.write_text(
                json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
            publications = [(staged, path), (staged_index, index_path)]
            backups = {}
            for position, (_new, destination) in enumerate(publications):
                if destination.exists():
                    backup = stage / f"previous-{position}"
                    os.link(destination, backup, follow_symlinks=False)
                    backups[destination] = backup
            published = []
            try:
                for new, destination in publications:
                    os.replace(new, destination)
                    published.append(destination)
            except BaseException:
                for destination in reversed(published):
                    if destination in backups:
                        os.replace(backups[destination], destination)
                    else:
                        destination.unlink()
                raise
        return {"index": str(index_path), "worksheets": len(workbook.worksheets)}
    finally:
        for sheet in workbook.worksheets:
            if sheet._writer is not None:
                if not sheet.closed:
                    with suppress(Exception):
                        sheet.close()
                if pathlib.Path(sheet._writer.out).exists():
                    sheet._writer.cleanup()
        workbook.close()
        template.close()


write_stream = write
