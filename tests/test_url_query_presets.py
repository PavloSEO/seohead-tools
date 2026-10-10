"""Ready-made filter presets and the filtered export of scan-url-query (issue #1007)."""

from __future__ import annotations

import csv
import json
import os

import pytest

from seohead import cli
from seohead.mcp import handlers
from seohead.storage import url_query
from tests.test_scan_url_query import _rows, _scan


@pytest.fixture(scope="module")
def scan_path(tmp_path_factory):
    return _scan(tmp_path_factory.mktemp("presets"))


def _urls(result):
    assert result["ok"] is True, result
    return sorted(row["url"] for row in result["rows"])


def _expected(predicate):
    return sorted(r["url"] for r in _rows() if predicate(r))


PREDICATES = {
    "status_4xx": lambda r: r["status_code"] is not None and 400 <= r["status_code"] < 500,
    "status_5xx": lambda r: r["status_code"] is not None and 500 <= r["status_code"] < 600,
    "no_response": lambda r: r["status_code"] is None,
    "title_missing": lambda r: r["title"] == "",
    "title_over_60": lambda r: len(r["title"]) > 60,
    "h1_missing": lambda r: r["h1"] == "",
    "noindex_meta": lambda r: "noindex" in r["meta_robots"].lower(),
}


@pytest.mark.parametrize("preset", sorted(PREDICATES))
def test_preset_matches_its_hand_written_predicate(scan_path, preset):
    got = handlers.scan_url_query(input_path=str(scan_path), preset=preset, limit=200)
    assert _urls(got) == _expected(PREDICATES[preset])
    assert got["preset"] == preset


def test_every_preset_equals_its_explicit_filters(scan_path):
    for preset, (_label, filters) in url_query.PRESETS.items():
        by_preset = handlers.scan_url_query(input_path=str(scan_path), preset=preset, limit=200)
        by_hand = handlers.scan_url_query(
            input_path=str(scan_path), filters=list(filters), limit=200
        )
        assert _urls(by_preset) == _urls(by_hand), preset


def test_preset_is_and_combined_with_user_filters(scan_path):
    got = handlers.scan_url_query(
        input_path=str(scan_path),
        preset="status_4xx",
        filters=[{"column": "status_code", "op": "eq", "value": 404}],
        limit=200,
    )
    assert _urls(got) == _expected(lambda r: r["status_code"] == 404)
    assert got["filters"][0]["column"] == "status_class"


def test_unknown_preset_is_refused(scan_path):
    got = handlers.scan_url_query(input_path=str(scan_path), preset="duplicate_titles")
    assert got["ok"] is False and got["reason_code"] == "unknown_preset"


def test_preset_does_not_exceed_the_filter_limit(scan_path):
    filters = [{"column": "status_code", "op": "gte", "value": 0}] * url_query.MAX_FILTERS
    got = handlers.scan_url_query(
        input_path=str(scan_path), preset="title_missing", filters=filters
    )
    assert got["ok"] is False and got["reason_code"] == "invalid_filter"


def _read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle, delimiter=";"))


def test_csv_export_writes_every_matching_row(scan_path, tmp_path):
    out = tmp_path / "missing-titles.csv"
    got = handlers.scan_url_query(
        input_path=str(scan_path),
        preset="title_missing",
        columns=["url", "title", "status_code"],
        export=str(out),
    )
    expected = _expected(PREDICATES["title_missing"])
    assert got["ok"] is True and got["rows"] == len(expected) > 0
    table = _read_csv(out)
    assert table[0] == ["url", "title", "status_code"]
    assert sorted(line[0] for line in table[1:]) == expected
    assert got["file"] == str(out)


def test_xlsx_export_reads_back_the_same_rows(scan_path, tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    out = tmp_path / "errors.xlsx"
    got = handlers.scan_url_query(
        input_path=str(scan_path),
        preset="status_5xx",
        columns=["url", "status_code"],
        export=str(out),
        export_format="xlsx",
    )
    assert got["ok"] is True
    sheet = load_workbook(out, read_only=True)["Pages"]
    body = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert sorted(body) == _expected(PREDICATES["status_5xx"])


def test_export_refuses_above_the_cap_and_writes_nothing(scan_path, tmp_path):
    out = tmp_path / "too-many.csv"
    got = handlers.scan_url_query(
        input_path=str(scan_path), preset="status_4xx", export=str(out), export_max_rows=1
    )
    assert got["ok"] is False and got["reason_code"] == "export_too_large"
    assert not out.exists()


def test_export_refuses_to_overwrite(scan_path, tmp_path):
    out = tmp_path / "exists.csv"
    out.write_text("keep me")
    got = handlers.scan_url_query(input_path=str(scan_path), preset="h1_missing", export=str(out))
    assert got["ok"] is False and got["reason_code"] == "output_exists"
    assert out.read_text() == "keep me"


def test_export_refuses_a_missing_directory_and_bad_format(scan_path, tmp_path):
    missing_dir = handlers.scan_url_query(
        input_path=str(scan_path), export=str(tmp_path / "nope" / "x.csv")
    )
    assert missing_dir["reason_code"] == "invalid_output"
    bad_format = handlers.scan_url_query(
        input_path=str(scan_path), export=str(tmp_path / "x.json"), export_format="json"
    )
    assert bad_format["reason_code"] == "invalid_export_format"
    assert not os.path.lexists(tmp_path / "x.json")


def test_export_leaves_the_scan_unchanged(scan_path, tmp_path):
    import hashlib

    before = hashlib.sha256(scan_path.read_bytes()).hexdigest()
    handlers.scan_url_query(
        input_path=str(scan_path), preset="title_missing", export=str(tmp_path / "a.csv")
    )
    assert hashlib.sha256(scan_path.read_bytes()).hexdigest() == before


def test_cli_preset_and_export_reach_the_handler(scan_path, tmp_path, capsys):
    out = tmp_path / "cli.csv"
    code = cli.main(
        [
            "scan-url-query",
            "--scan",
            str(scan_path),
            "--preset",
            "title_missing",
            "--columns",
            "url,title",
            "--export",
            str(out),
        ]
    )
    result = json.loads(capsys.readouterr().out)
    assert code == 0 and result["ok"] is True and result["preset"] == "title_missing"
    assert len(_read_csv(out)) - 1 == result["rows"] == len(_expected(PREDICATES["title_missing"]))
