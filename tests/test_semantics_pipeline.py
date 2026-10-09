from __future__ import annotations

from pathlib import Path
from typing import ClassVar

import pytest
from openpyxl import load_workbook

from seohead.semantics.stages import (
    clean,
    collect,
    exact,
    excel,
    graph,
    report,
    sitematch,
    synonyms,
)

pytestmark = pytest.mark.usefixtures("semantics_offline")


def test_collect_retries_a_transient_non_quota_failure_instead_of_poisoning_the_seed(
    sem_store, sem_cfg, monkeypatch
):
    class FakeWordstat:
        calls = 0

        def __init__(self, **_kwargs):
            pass

        def expand(self, phrase, *, limit, regions):
            type(self).calls += 1
            if type(self).calls == 1:
                raise RuntimeError("topRequests 502: temporary upstream error")
            assert phrase == "краулер сайта"
            return {"краулер сайта онлайн": 80}, {
                "results": 1,
                "associations": 0,
                "origin": {"краулер сайта онлайн": "results"},
            }

    monkeypatch.setattr(collect, "Wordstat", FakeWordstat)
    first = collect.run(sem_store, sem_cfg, seeds=["краулер сайта"], max_seeds=1)
    assert first["quota_hit"] is False
    assert first["expanded"] == 0
    assert sem_store.is_expanded("краулер сайта") is False

    second = collect.run(sem_store, sem_cfg, seeds=["краулер сайта"], max_seeds=1)
    assert second["quota_hit"] is False
    assert second["expanded"] == 1
    assert FakeWordstat.calls == 2
    assert (
        sem_store.db.execute(
            "SELECT base FROM phrases WHERE norm='краулер сайта онлайн'"
        ).fetchone()[0]
        == 80
    )


def test_exact_uses_overal_for_wordform_frequency_and_preserves_zero_and_missing(
    sem_store, sem_cfg, monkeypatch
):
    sem_store.upsert("seo аудит сайта", base=100)
    sem_store.upsert("seo аудит цена", base=90)
    sem_store.set_fields("seo аудит сайта", status="kept", lemma_group=1)
    sem_store.set_fields("seo аудит цена", status="kept", lemma_group=2)
    sem_store.commit()

    class FakeClient:
        submitted: ClassVar[list] = []

        def set_task(self, tool, payload):
            type(self).submitted.append((tool, payload))
            return {"task_id": 314, "cost": 2}

        def wait(self, task_id):
            assert task_id == 314
            return {
                "code": "TASK_RESULT",
                "result": {
                    "data": {
                        "result": {
                            "seo аудит сайта": {
                                "225": {"base": 100, "quoted": 17, "overal": 0, "exact": 3}
                            }
                        }
                    }
                },
            }

    monkeypatch.setattr(exact, "ArsenkinClient", FakeClient)
    result = exact.run(sem_store, sem_cfg)

    assert result["snapped"] == 1
    assert FakeClient.submitted == [
        (
            "wordstat",
            {
                "type": 1,
                "regions": [225],
                "ws": ["base", "overal", "quoted"],
                "queries": ["seo аудит сайта", "seo аудит цена"],
            },
        )
    ]
    measured = sem_store.db.execute(
        "SELECT quoted, exact FROM phrases WHERE norm='seo аудит сайта'"
    ).fetchone()
    absent = sem_store.db.execute(
        "SELECT quoted, exact FROM phrases WHERE norm='seo аудит цена'"
    ).fetchone()
    assert tuple(measured) == (17, 0)
    assert tuple(absent) == (None, None)


def test_synonyms_uses_the_toolkit_signature_and_caches_rejected_candidates(
    sem_store, sem_cfg, monkeypatch, tmp_path
):
    (tmp_path / "synonyms.txt").write_text("аудит сайта бесплатно\n", encoding="utf-8")

    class FakeWordstat:
        calls: ClassVar[list] = []

        def __init__(self, **_kwargs):
            pass

        def top(self, phrase, *, limit=300, regions=()):
            type(self).calls.append((phrase, limit, tuple(regions)))
            return {"totalCount": 3}

    monkeypatch.setattr(synonyms, "Wordstat", FakeWordstat)
    first = synonyms.run(sem_store, sem_cfg)
    second = synonyms.run(sem_store, sem_cfg)

    assert first["added"] == 0 and first["checked"] == 1
    assert second["added"] == 0
    assert FakeWordstat.calls == [("аудит сайта бесплатно", 1, ("225",))]
    row = sem_store.db.execute(
        "SELECT base, status, reason_stage FROM phrases WHERE norm='аудит сайта бесплатно'"
    ).fetchone()
    assert tuple(row) == (3, "dropped", "synonyms")


def test_local_clean_graph_sitematch_report_and_excel_keep_evidence_and_emit_artifacts(
    sem_store, sem_cfg, tmp_path
):
    sem_store.upsert("seo аудит сайта", base=100, src="webmaster")
    sem_store.set_fields("seo аудит сайта", impr=37, pos_y=4, pos_g=8)
    sem_store.upsert("как провести seo аудит", base=60)
    sem_store.upsert("краулер сайта онлайн", base=50)
    sem_store.upsert("скачать бесплатно seo аудит", base=200)
    sem_store.commit()

    cleaned = clean.run(sem_store, sem_cfg)
    assert cleaned["changed"] == 4
    assert tuple(
        sem_store.db.execute(
            "SELECT status, route, impr, pos_y, pos_g FROM phrases WHERE norm='seo аудит сайта'"
        ).fetchone()
    ) == ("kept", "commercial", 37, 4, 8)
    assert tuple(
        sem_store.db.execute(
            "SELECT status, route FROM phrases WHERE norm='как провести seo аудит'"
        ).fetchone()
    ) == ("kept", "info")
    assert (
        sem_store.db.execute(
            "SELECT status FROM phrases WHERE norm='скачать бесплатно seo аудит'"
        ).fetchone()[0]
        == "dropped"
    )

    (tmp_path / "catalog.txt").write_text("аудит\n", encoding="utf-8")
    assert sitematch.run(sem_store, sem_cfg)["off_catalog"] == 1
    assert (
        sem_store.db.execute(
            "SELECT status FROM phrases WHERE norm='краулер сайта онлайн'"
        ).fetchone()[0]
        == "review"
    )

    graph_result = graph.run(sem_store, sem_cfg, tmp_path)
    assert graph_result["homonyms"] == 0
    assert (tmp_path / "anomalies.md").is_file()

    cid = sem_store.upsert_cluster("seo аудит сайта", source="fixture")
    sem_store.assign_cluster("seo аудит сайта", cid)
    sem_store.commit()
    assert report.run(sem_store, sem_cfg, tmp_path)["clusters"] == 1
    workbook_result = excel.run(sem_store, sem_cfg, tmp_path)
    assert Path(workbook_result["file"]).is_file()
    assert "seo аудит сайта" in (tmp_path / "report.md").read_text(encoding="utf-8")
    book = load_workbook(workbook_result["file"], read_only=True)
    assert {"Summary", "Core", "Clusters", "Rejected"} <= set(book.sheetnames)


def test_graph_does_not_demote_observed_search_demand(sem_store, sem_cfg, tmp_path, monkeypatch):
    phrase = "synthetic off topic"
    sem_store.upsert(phrase, base=0)
    sem_store.set_fields(phrase, pos_g=3, status="kept")
    monkeypatch.setattr(graph, "detect_homonyms", lambda phrases: {"synthetic": [{phrase}]})
    graph.run(sem_store, sem_cfg, tmp_path)
    assert sem_store.phrases()[0]["status"] == "kept"
