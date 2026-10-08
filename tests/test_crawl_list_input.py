"""Explicit list-mode URL file import is local and deterministic."""

import csv
from pathlib import Path

import pytest

from seohead.crawl.list_input import UrlListError, iter_url_list, read_url_list


def test_text_input_scans_urls_and_preserves_their_order(tmp_path):
    source = tmp_path / "urls.txt"
    source.write_text(
        "note https://example.test/one, then https://example.test/two?x=1\n",
        encoding="utf-8",
    )

    assert read_url_list(source) == ["https://example.test/one", "https://example.test/two?x=1"]


def test_csv_input_scans_every_column_in_row_order(tmp_path):
    source = tmp_path / "redirects.csv"
    with source.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["old", "note"])
        writer.writerow(["https://example.test/old", "to https://example.test/new"])

    assert read_url_list(source) == ["https://example.test/old", "https://example.test/new"]


def test_spreadsheet_input_scans_cells_in_workbook_order(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "urls.xlsx"
    workbook = Workbook()
    workbook.active.append(["https://example.test/one", "other"])
    workbook.create_sheet("second").append(["https://example.test/two"])
    workbook.save(source)

    assert read_url_list(source) == ["https://example.test/one", "https://example.test/two"]


def test_xml_input_scans_attribute_and_text_values_in_document_order(tmp_path):
    source = tmp_path / "urls.xml"
    source.write_text(
        '<root href="https://example.test/one"><loc>https://example.test/two</loc></root>',
        encoding="utf-8",
    )

    assert read_url_list(source) == ["https://example.test/one", "https://example.test/two"]


def test_unsupported_or_empty_input_is_refused_by_name(tmp_path):
    unsupported = tmp_path / "urls.json"
    unsupported.write_text("https://example.test/ignored", encoding="utf-8")
    empty = tmp_path / "urls.txt"
    empty.write_text("no URLs here", encoding="utf-8")

    with pytest.raises(UrlListError, match="unsupported"):
        read_url_list(unsupported)
    with pytest.raises(UrlListError, match="no absolute"):
        read_url_list(empty)


@pytest.mark.parametrize("suffix", [".txt", ".csv"])
def test_streaming_text_inputs_close_their_handle_when_stopped_early(tmp_path, monkeypatch, suffix):
    source = tmp_path / ("urls" + suffix)
    source.write_text("https://example.test/one\nhttps://example.test/two\n")
    original = Path.open
    opened = []

    def traced(path, *args, **kwargs):
        stream = original(path, *args, **kwargs)
        if path == source:
            opened.append(stream)
        return stream

    monkeypatch.setattr(Path, "open", traced)
    iterator = iter_url_list(source)
    assert next(iterator) == "https://example.test/one"
    assert len(opened) == 1 and not opened[0].closed
    iterator.close()
    assert opened[0].closed


def test_streaming_spreadsheet_closes_workbook_when_stopped_early(tmp_path, monkeypatch):
    import openpyxl

    source = tmp_path / "urls.xlsx"
    source.touch()
    closed = []

    class Worksheet:
        def iter_rows(self, **_kwargs):
            yield ("https://example.test/one",)
            raise AssertionError("the second row must remain unread")

    class Workbook:
        worksheets = (Worksheet(),)

        def close(self):
            closed.append(True)

    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_a, **_kw: Workbook())
    iterator = iter_url_list(source)
    assert next(iterator) == "https://example.test/one"
    iterator.close()
    assert closed == [True]


@pytest.mark.parametrize("suffix", [".txt", ".csv"])
def test_oversized_physical_text_record_is_refused_without_loading_the_file(
    tmp_path, monkeypatch, suffix
):
    from seohead.crawl import list_input

    monkeypatch.setattr(list_input, "_MAX_TEXT_RECORD_CHARS", 32)
    source = tmp_path / ("urls" + suffix)
    source.write_text("https://example.test/" + "x" * 100)
    with pytest.raises(UrlListError, match="text record exceeds"):
        next(iter_url_list(source))
