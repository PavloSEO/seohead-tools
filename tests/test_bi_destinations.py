from __future__ import annotations

import json
import re
from hashlib import sha256
from pathlib import Path

import pytest

from seohead import cli
from seohead.reports.bi import export_bi
from seohead.reports.bi_destinations import (
    BIDestinationCommitUncertain,
    BIDestinationError,
    GoogleBigQueryClient,
    GoogleSheetsClient,
    apply_with_client,
    bigquery_plan,
    destination_preview,
    filter_package,
    register_host_client,
    resolve_host_client,
    sheets_plan,
)
from seohead.servers import handlers


def _audit():
    return {
        "schema": "seohead.site-audit/1",
        "url": "https://example.test/",
        "site": {},
        "pages": [{"url": "https://example.test/", "status_code": 200}],
        "findings": [],
        "summary": {"pages_checked": 1, "findings_total": 0, "tools_run": [], "tools_failed": []},
    }


def test_offline_destination_plans_reconcile_the_complete_local_package(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    sheets = sheets_plan(package)
    assert sheets["network"] is False and sheets["apply"] is False
    assert sheets["state"] == "ready"
    assert {row["worksheet"] for row in sheets["worksheets"]} >= {"pages", "cohorts"}
    bigquery = bigquery_plan(package, dataset="synthetic_reporting")
    assert bigquery["billing_required"] is True
    assert all(row["table"].startswith("seohead_") for row in bigquery["tables"])
    assert sheets_plan(package, max_cells=1)["state"] == "unavailable"


def test_injected_destination_client_is_explicit_transactional_and_streamed(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)

    class Client:
        def __init__(self):
            self.calls = []

        def authorize_target(self, target):
            self.calls.append(("authorize", target))
            return True

        def begin(self, **kwargs):
            self.calls.append(("begin", kwargs))
            return "tx"

        def write(self, transaction, dataset, rows):
            self.calls.append(("write", transaction, dataset, len(rows)))

        def commit(self, transaction):
            self.calls.append(("commit", transaction))

    client = Client()
    with pytest.raises(BIDestinationError, match="apply=True"):
        apply_with_client(package, target="synthetic", operation="replace", client=client)
    result = apply_with_client(
        package, target="synthetic", operation="replace", client=client, apply=True
    )
    assert result["rows"]["pages"] == 1
    assert client.calls[0] == ("authorize", "synthetic")
    assert client.calls[-1] == ("commit", "tx")


def test_filtered_bi_export_is_exact_and_partitioned(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    filtered = filter_package(
        package,
        dataset="cohorts",
        out_dir=tmp_path / "filtered",
        where={"cohort_id": ["observed_status", "crawl_relative_depth"]},
        columns=["run_id", "cohort_id", "value_label", "state"],
        max_rows_per_file=1,
    )
    assert filtered["row_count"] == 2
    assert [part["rows"] for part in filtered["partitions"]] == [1, 1]
    assert (tmp_path / "filtered" / "manifest.json").is_file()


def test_destination_preflight_streams_partitions_and_accepts_selected_projection(
    tmp_path, monkeypatch
):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    filtered = filter_package(
        package,
        dataset="cohorts",
        out_dir=tmp_path / "filtered",
        columns=["run_id", "cohort_id", "state"],
    )
    assert filtered["row_count"] > 0
    monkeypatch.setattr(
        Path, "read_bytes", lambda *_args: pytest.fail("must stream partition bytes")
    )
    plan = sheets_plan(tmp_path / "filtered")
    assert [worksheet["worksheet"] for worksheet in plan["worksheets"]] == ["cohorts"]
    assert plan["worksheets"][0]["columns"] == 3


def test_selected_bi_view_streams_to_excel_with_numbered_sheet_splitting(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    filtered = filter_package(
        package,
        dataset="cohorts",
        out_dir=tmp_path / "filtered",
        columns=["run_id", "cohort_id", "state"],
    )
    from seohead.reports.bi_destinations import export_bi_xlsx

    result = export_bi_xlsx(
        tmp_path / "filtered",
        dataset="cohorts",
        out=tmp_path / "cohorts.xlsx",
        max_rows_per_sheet=2,
    )
    assert result["rows"] == filtered["row_count"]
    assert result["worksheets"] == (filtered["row_count"] + 1) // 2
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "cohorts.xlsx", read_only=True)
    try:
        assert workbook.sheetnames[0] == "cohorts-0001"
        assert next(iter(workbook["cohorts-0001"].values)) == ("run_id", "cohort_id", "state")
    finally:
        workbook.close()


def test_selected_projection_preview_uses_its_mapping_from_a_full_target_config(
    tmp_path, monkeypatch
):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    filter_package(
        package,
        dataset="cohorts",
        out_dir=tmp_path / "filtered",
        columns=["run_id", "cohort_id", "state"],
    )
    config = tmp_path / "bi-destinations.json"
    config.write_text(
        json.dumps(
            {
                "sheets": {
                    "targets": {
                        "reporting": {
                            "enabled": True,
                            "kind": "google_sheets_service_account",
                            "spreadsheet_id": "1Qs8BdfxZXALh6vX4zrE7ZyGnR3h5k",
                            "worksheets": _worksheet_mapping(package),
                        }
                    }
                }
            }
        )
    )
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))
    preview = destination_preview(
        tmp_path / "filtered", target="reporting", destination="sheets", operation="replace"
    )
    assert set(preview["resolved_target"]["worksheets"]) == {"cohorts"}


def test_destination_preflight_rejects_manifest_type_or_csv_header_drift(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["datasets"]["pages"]["fields"][0]["type"] = "integer"
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(BIDestinationError, match="fields do not match"):
        sheets_plan(package)

    package = tmp_path / "header-package"
    export_bi(audit=_audit(), out_dir=package)
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    part = manifest["datasets"]["pages"]["partitions"][0]
    path = package / part["path"]
    content = path.read_text()
    path.write_text(content.replace("run_id", "wrong_header", 1))
    raw = path.read_bytes()
    old_bytes = part["bytes"]
    part["bytes"] = len(raw)
    part["sha256"] = sha256(raw).hexdigest()
    manifest["datasets"]["pages"]["bytes"] += len(raw) - old_bytes
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(BIDestinationError, match="header"):
        sheets_plan(package)


def test_shared_handler_uses_exact_host_allowlist_and_registered_client(tmp_path, monkeypatch):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    config = tmp_path / "bi-destinations.json"
    config.write_text(json.dumps({"sheets": {"targets": {"synthetic": {"enabled": True}}}}))
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))
    with pytest.raises(ValueError, match="authorized client"):
        handlers.bi_destination_apply(
            package=str(package),
            target="synthetic",
            destination="sheets",
            operation="replace",
            apply=True,
        )

    class Client:
        def authorize_target(self, target):
            return target == "synthetic"

        def begin(self, **_kwargs):
            return "tx"

        def write(self, *_args):
            pass

        def commit(self, _transaction):
            pass

    result = handlers.bi_destination_apply(
        package=str(package),
        target="synthetic",
        destination="sheets",
        operation="replace",
        apply=True,
        client=Client(),
    )
    assert result["destination"] == "sheets" and result["rows"]["pages"] == 1

    register_host_client("sheets", "synthetic", Client())
    hosted = handlers.bi_destination_apply(
        package=str(package),
        target="synthetic",
        destination="sheets",
        operation="replace",
        apply=True,
    )
    assert hosted["rows"]["pages"] == 1


def test_destination_preview_shows_resolved_target_and_operation_without_auth_or_transport(
    tmp_path, monkeypatch
):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    worksheets = _worksheet_mapping(package)
    config = tmp_path / "bi-destinations.json"
    config.write_text(
        json.dumps(
            {
                "sheets": {
                    "targets": {
                        "reporting": {
                            "enabled": True,
                            "kind": "google_sheets_service_account",
                            "spreadsheet_id": "1Qs8BdfxZXALh6vX4zrE7ZyGnR3h5k",
                            "worksheets": worksheets,
                        }
                    }
                }
            }
        )
    )
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))
    auth_calls = []
    monkeypatch.setattr(
        "seohead.data_sources.gsc.service_account_access_token",
        lambda *_args: auth_calls.append(True) or pytest.fail("preview must not authenticate"),
    )
    result = handlers.bi_destination_apply(
        package=str(package),
        target="reporting",
        destination="sheets",
        operation="replace",
    )
    assert result["state"] == "ready_to_apply"
    assert result["target"] == "reporting" and result["operation"] == "replace"
    assert result["datasets"]["pages"]["rows"] == 1
    assert result["resolved_target"] == {
        "kind": "google_sheets_service_account",
        "spreadsheet_id": "1Qs8BdfxZXALh6vX4zrE7ZyGnR3h5k",
        "worksheets": worksheets,
    }
    assert auth_calls == []


def test_bigquery_preview_discloses_resolved_project_dataset_and_table_mapping(
    tmp_path, monkeypatch
):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    tables = _table_mapping(package)
    config = tmp_path / "bi-destinations.json"
    config.write_text(
        json.dumps(
            {
                "bigquery": {
                    "targets": {
                        "warehouse": {
                            "enabled": True,
                            "kind": "google_bigquery_service_account",
                            "project_id": "reporting-project",
                            "dataset_id": "seohead_reporting",
                            "location": "EU",
                            "cost_authorized": False,
                            "tables": tables,
                        }
                    }
                }
            }
        )
    )
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))
    result = handlers.bi_destination_apply(
        package=str(package),
        target="warehouse",
        destination="bigquery",
        operation="replace",
    )
    assert result["resolved_target"] == {
        "kind": "google_bigquery_service_account",
        "project_id": "reporting-project",
        "dataset_id": "seohead_reporting",
        "location": "EU",
        "cost_authorized": False,
        "tables": tables,
    }


def _worksheet_mapping(package):
    manifest = json.loads((package / "manifest.json").read_text())
    return {
        name: {"worksheet_id": index + 2, "worksheet_title": name}
        for index, name in enumerate(manifest["datasets"])
    }


def _table_mapping(package):
    manifest = json.loads((package / "manifest.json").read_text())
    return {name: {"table_id": f"seohead_{name}_v1"} for name in manifest["datasets"]}


def test_google_sheets_replace_uses_one_atomic_swap_and_preserves_target_sheet_ids(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    requests = []
    staged_values = {}
    worksheets = _worksheet_mapping(package)

    def fetch(request):
        requests.append(request)
        if request["method"] == "GET":
            if "/values/" not in request["url"]:
                return {
                    "sheets": [
                        {
                            "properties": {
                                "sheetId": item["worksheet_id"],
                                "title": item["worksheet_title"],
                                "gridProperties": {"rowCount": 1, "columnCount": 1},
                            }
                        }
                        for item in worksheets.values()
                    ]
                }
            title = request["url"].split("/values/'", 1)[1].split("'!", 1)[0]
            return {"values": staged_values[title]}
        if (
            request["url"].endswith(":batchUpdate")
            and "requests" in request["body"]
            and "addSheet" in request["body"]["requests"][0]
        ):
            adds = request["body"]["requests"]
            return {
                "replies": [
                    {
                        "addSheet": {
                            "properties": {
                                "sheetId": 77 + index,
                                "title": item["addSheet"]["properties"]["title"],
                            }
                        }
                    }
                    for index, item in enumerate(adds)
                ]
            }
        if request["url"].endswith("/values:batchUpdate"):
            item = request["body"]["data"][0]
            title = item["range"].split("'", 2)[1]
            staged_values[title] = item["values"]
            return {"totalUpdatedRows": len(staged_values[title])}
        return {}

    client = GoogleSheetsClient(
        "synthetic",
        "sheet-id",
        worksheets,
        token_supplier=lambda scope: "token",
        fetcher=fetch,
    )
    result = apply_with_client(
        package, target="synthetic", operation="replace", client=client, apply=True
    )
    assert result["state"] == "committed"
    assert result["row_conservation"] == "verified"
    commit = requests[-1]["body"]["requests"]
    assert all(
        "updateCells" in request or "copyPaste" in request or "deleteSheet" in request
        for request in commit
    )
    assert {
        request["copyPaste"]["destination"]["sheetId"]
        for request in commit
        if "copyPaste" in request
    } == {item["worksheet_id"] for item in worksheets.values()}
    assert any(
        request["updateCells"].get("fields") == "userEnteredValue"
        for request in commit
        if "updateCells" in request
    )
    assert requests[1]["authorization"] == "Bearer token"


def test_google_sheets_rejects_partial_mapping_before_creating_a_stage(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    called = []
    client = GoogleSheetsClient(
        "synthetic",
        "sheet-id",
        {"pages": {"worksheet_id": 2, "worksheet_title": "pages"}},
        token_supplier=lambda scope: "token",
        fetcher=lambda request: called.append(request) or {},
    )
    with pytest.raises(BIDestinationError, match="exactly once"):
        apply_with_client(
            package, target="synthetic", operation="replace", client=client, apply=True
        )
    assert called == []


def test_google_bigquery_stages_chunks_and_publishes_each_table_with_mocked_rest(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    requests = []
    expected = {
        name: item["row_count"]
        for name, item in json.loads((package / "manifest.json").read_text())["datasets"].items()
    }
    job_rows = {}

    def fetch(request):
        requests.append(request)
        if request["method"] == "POST" and "upload/bigquery" in request["url"]:
            job_id = re.search(rb'"jobId":"([^"]+)"', request["body"]).group(1).decode()
            media = request["body"].split(b"Content-Type: application/octet-stream\r\n\r\n", 1)[1]
            media = media.rsplit(b"\r\n--", 1)[0]
            job_rows[job_id] = media.count(b"\n")
            return {"jobReference": {"jobId": job_id}}
        if request["method"] == "POST" and request["url"].endswith("/jobs"):
            return {"jobReference": {"jobId": request["body"]["jobReference"]["jobId"]}}
        if request["method"] == "GET" and "/jobs/" in request["url"]:
            job_id = request["url"].split("/jobs/", 1)[1].split("?", 1)[0]
            return {
                "status": {"state": "DONE"},
                "statistics": {"load": {"outputRows": str(job_rows.get(job_id, 0))}},
            }
        if (
            request["method"] == "GET"
            and "/datasets/dataset_id" in request["url"]
            and "/tables/" not in request["url"]
        ):
            return {"datasetReference": {"datasetId": "dataset_id"}}
        if request["method"] == "GET" and "/tables/" in request["url"]:
            stage = request["url"].rsplit("/", 1)[1]
            name = re.match(r"_seohead_stage_(.+)_[0-9a-f]{16}$", stage).group(1)
            return {"numRows": str(expected[name])}
        if request["method"] == "DELETE":
            return {}
        raise AssertionError(request)

    client = GoogleBigQueryClient(
        "synthetic",
        "project-id",
        "dataset_id",
        _table_mapping(package),
        location="EU",
        cost_authorized=True,
        token_supplier=lambda scope: "token",
        fetcher=fetch,
    )
    result = apply_with_client(
        package, target="synthetic", operation="replace", client=client, apply=True
    )
    assert result["publication"] == "atomic_per_bigquery_table_not_across_datasets"
    uploads = [request for request in requests if "upload/bigquery" in request["url"]]
    assert uploads and all(isinstance(request["body"], bytes) for request in uploads)
    copies = [
        request
        for request in requests
        if request["method"] == "POST" and request["url"].endswith("/jobs")
    ]
    assert copies and all(
        request["body"]["configuration"]["copy"]["writeDisposition"] == "WRITE_TRUNCATE"
        for request in copies
    )


def test_fresh_host_resolution_creates_google_clients_without_registration(tmp_path, monkeypatch):
    config = tmp_path / "bi-destinations.json"
    config.write_text(
        json.dumps(
            {
                "sheets": {
                    "targets": {
                        "reporting": {
                            "enabled": True,
                            "kind": "google_sheets_service_account",
                            "spreadsheet_id": "spreadsheet-id",
                            "worksheets": {},
                        }
                    }
                },
                "bigquery": {
                    "targets": {
                        "warehouse": {
                            "enabled": True,
                            "kind": "google_bigquery_service_account",
                            "project_id": "project-id",
                            "dataset_id": "dataset_id",
                            "tables": {},
                            "cost_authorized": True,
                        }
                    }
                },
            }
        )
    )
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))
    assert isinstance(resolve_host_client("sheets", "reporting"), GoogleSheetsClient)
    assert isinstance(resolve_host_client("bigquery", "warehouse"), GoogleBigQueryClient)


def test_bigquery_cost_denial_calls_neither_auth_nor_transport(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    calls = []
    client = GoogleBigQueryClient(
        "warehouse",
        "project-id",
        "dataset_id",
        _table_mapping(package),
        token_supplier=lambda scope: calls.append(("auth", scope)) or "token",
        fetcher=lambda request: calls.append(("transport", request)) or {},
    )
    with pytest.raises(BIDestinationError, match="cost_authorized"):
        apply_with_client(
            package, target="warehouse", operation="replace", client=client, apply=True
        )
    assert calls == []


def test_host_bigquery_apply_rejects_missing_cost_authorization_before_auth(tmp_path, monkeypatch):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    config = tmp_path / "bi-destinations.json"
    config.write_text(
        json.dumps(
            {
                "bigquery": {
                    "targets": {
                        "warehouse": {
                            "enabled": True,
                            "kind": "google_bigquery_service_account",
                            "project_id": "project-id",
                            "dataset_id": "dataset_id",
                            "tables": _table_mapping(package),
                        }
                    }
                }
            }
        )
    )
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))
    auth_calls = []
    monkeypatch.setattr(
        "seohead.data_sources.gsc.service_account_access_token",
        lambda *_args: auth_calls.append(True) or pytest.fail("auth must not run"),
    )
    with pytest.raises(ValueError, match="cost_authorized"):
        handlers.bi_destination_apply(
            package=str(package),
            target="warehouse",
            destination="bigquery",
            operation="replace",
            apply=True,
        )
    assert auth_calls == []


def test_cli_and_mcp_resolve_the_same_configured_sheets_client_offline(tmp_path, monkeypatch):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    config = tmp_path / "bi-destinations.json"
    config.write_text(
        json.dumps(
            {
                "sheets": {
                    "targets": {
                        "reporting": {
                            "enabled": True,
                            "kind": "google_sheets_service_account",
                            "spreadsheet_id": "spreadsheet-id",
                            "worksheets": _worksheet_mapping(package),
                        }
                    }
                }
            }
        )
    )
    monkeypatch.setenv("SEOHEAD_BI_DESTINATIONS_FILE", str(config))

    worksheets = _worksheet_mapping(package)

    def request(self, method, url, body=None, *, headers=None):
        if method == "GET":
            if "/values/" not in url:
                return {
                    "sheets": [
                        {
                            "properties": {
                                "sheetId": item["worksheet_id"],
                                "title": item["worksheet_title"],
                                "gridProperties": {"rowCount": 1, "columnCount": 1},
                            }
                        }
                        for item in worksheets.values()
                    ]
                }
            return {"values": body or [["placeholder"]]}
        if "addSheet" in str(body):
            return {
                "replies": [
                    {
                        "addSheet": {
                            "properties": {
                                "sheetId": 100 + index,
                                "title": item["addSheet"]["properties"]["title"],
                            }
                        }
                    }
                    for index, item in enumerate(body["requests"])
                ]
            }
        if url.endswith("/values:batchUpdate"):
            # The next GET only needs a complete exact range in this transport
            # seam; this patch avoids authentication and all external traffic.
            request.last_values = body["data"][0]["values"]
            return {"totalUpdatedRows": len(body["data"][0]["values"])}
        return {}

    original_request = GoogleSheetsClient._request

    def request_with_readback(self, method, url, body=None, *, headers=None, **_kwargs):
        if method == "GET" and "/values/" in url:
            return {"values": request.last_values}
        return request(self, method, url, body, headers=headers)

    request.last_values = []
    monkeypatch.setattr(GoogleSheetsClient, "_request", request_with_readback)
    assert (
        cli.main(
            [
                "bi-destination-apply",
                "--package",
                str(package),
                "--target",
                "reporting",
                "--destination",
                "sheets",
                "--apply",
            ]
        )
        == 0
    )
    from seohead.servers.mcp_server import build_server

    tool = build_server()._tool_manager.get_tool("seo_bi_destination_apply")
    result = tool.fn(
        package=str(package),
        target="reporting",
        destination="sheets",
        operation="replace",
        apply=True,
    )
    assert result["state"] == "committed"
    monkeypatch.setattr(GoogleSheetsClient, "_request", original_request)


def _restartable_sheets_fetcher(worksheets):
    """Synthetic stateful REST seam; no Google request or credential is used."""
    remote = {
        "sheets": {
            item["worksheet_id"]: {
                "sheetId": item["worksheet_id"],
                "title": item["worksheet_title"],
                "gridProperties": {"rowCount": 2, "columnCount": 100},
            }
            for item in worksheets.values()
        },
        "values": {},
    }

    def fetch(request):
        url, body = request["url"], request["body"]
        if request["method"] == "GET" and "/values/" not in url:
            return {"sheets": [{"properties": value} for value in remote["sheets"].values()]}
        if request["method"] == "GET" and "/values/" in url:
            return {"values": remote["values"].get(url.split("/values/", 1)[1], [])}
        if url.endswith("/values:batchUpdate"):
            values = body["data"][0]["values"]
            remote["values"][body["data"][0]["range"]] = values
            return {"totalUpdatedRows": len(values)}
        requests = body["requests"]
        if requests and "addSheet" in requests[0]:
            replies = []
            for index, item in enumerate(requests, start=100):
                props = item["addSheet"]["properties"]
                remote["sheets"][index] = {
                    "sheetId": index,
                    "title": props["title"],
                    "gridProperties": props["gridProperties"],
                }
                replies.append({"addSheet": {"properties": remote["sheets"][index]}})
            return {"replies": replies}
        for item in requests:
            if "deleteSheet" in item:
                remote["sheets"].pop(item["deleteSheet"]["sheetId"], None)
        return {}

    return fetch


def test_sheets_restart_reconciles_an_uncertain_stage_range_before_resuming(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)
    worksheets = _worksheet_mapping(package)
    fetch = _restartable_sheets_fetcher(worksheets)
    first = GoogleSheetsClient(
        "synthetic", "sheet-id", worksheets, token_supplier=lambda _scope: "token", fetcher=fetch
    )
    original_write = first.write
    calls = 0

    def interrupted_write(transaction, dataset, rows):
        nonlocal calls
        original_write(transaction, dataset, rows)
        calls += 1
        if calls == 2:
            raise KeyboardInterrupt

    first.write = interrupted_write
    with pytest.raises(KeyboardInterrupt):
        apply_with_client(
            package, target="synthetic", operation="replace", client=first, apply=True
        )
    assert list((package / ".seohead-destination-state").glob("*.json"))

    resumed = GoogleSheetsClient(
        "synthetic", "sheet-id", worksheets, token_supplier=lambda _scope: "token", fetcher=fetch
    )
    result = apply_with_client(
        package,
        target="synthetic",
        operation="replace",
        client=resumed,
        apply=True,
        reconcile=True,
    )
    assert result["state"] == "committed"
    assert result["rows"]["pages"] == result["input_rows"]["pages"] == 1
    assert result["failed_rows"]["pages"] == 0


def test_uncertain_write_never_replays_without_explicit_reconciliation(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)

    class Client:
        destination = "synthetic"

        def __init__(self):
            self.calls = 0

        def authorize_target(self, target):
            return target == "synthetic"

        def begin(self, **_kwargs):
            return {}

        def write(self, _transaction, _dataset, _rows):
            self.calls += 1
            if self.calls == 2:
                raise BIDestinationCommitUncertain("synthetic write response lost")

        def commit(self, _transaction):
            pytest.fail("uncertain write must stop before commit")

    client = Client()
    first = apply_with_client(
        package, target="synthetic", operation="replace", client=client, apply=True
    )
    assert first["state"] == "reconciliation_required"
    assert first["rows"]["pages"] == 0
    calls = client.calls
    second = apply_with_client(
        package, target="synthetic", operation="replace", client=client, apply=True
    )
    assert second["state"] == "reconciliation_required"
    assert client.calls == calls


def test_failed_write_returns_complete_per_dataset_accounting(tmp_path):
    package = tmp_path / "package"
    export_bi(audit=_audit(), out_dir=package)

    class Client:
        destination = "synthetic-failure"

        def __init__(self):
            self.calls = 0
            self.aborted = False

        def authorize_target(self, _target):
            return True

        def begin(self, **_kwargs):
            return {}

        def write(self, _transaction, _dataset, _rows):
            self.calls += 1
            if self.calls == 2:
                raise BIDestinationError("synthetic rejected bounded row")

        def abort(self, _transaction):
            self.aborted = True

        def commit(self, _transaction):
            pytest.fail("failed write must stop before commit")

    client = Client()
    result = apply_with_client(
        package, target="synthetic", operation="replace", client=client, apply=True
    )
    assert result["state"] == "failed" and client.aborted
    assert result["input_rows"]["cohorts"] == result["failed_rows"]["cohorts"] == 5
    assert result["rows"]["cohorts"] == 0
    assert set(result["input_rows"]) == set(result["rows"]) == set(result["failed_rows"])


def test_bigquery_duplicate_job_is_reconciled_by_its_deterministic_identity():
    calls = []

    def fetch(request):
        calls.append(request)
        if request["method"] == "POST":
            return {"error": {"code": 409, "message": "Already Exists"}}
        assert request["method"] == "GET" and request["url"].endswith("/jobs/known-job")
        return {"jobReference": {"jobId": "known-job"}, "status": {"state": "DONE"}}

    client = GoogleBigQueryClient(
        "synthetic",
        "project-id",
        "dataset_id",
        {},
        cost_authorized=True,
        token_supplier=lambda _scope: "token",
        fetcher=fetch,
    )
    result = client._job("known-job", {"copy": {}})
    assert result["status"]["state"] == "DONE"
    assert [request["method"] for request in calls] == ["POST", "GET"]
