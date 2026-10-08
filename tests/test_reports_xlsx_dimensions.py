"""Streaming workbook dimensions remain usable by read-only consumers."""

from __future__ import annotations

import gc
import hashlib
import json
import os
import sys

import pytest

from seohead.reports import xlsx
from tests.test_scan_artifact_office import frozen_office_clock as frozen_office_clock


def _document(count):
    return {
        "domain": "example.test",
        "findings": [
            {
                "severity": "warning",
                "url": f"https://example.test/{index}",
                "client_title": "Missing title",
            }
            for index in range(count)
        ],
        "pages": [
            {"url": f"https://example.test/{index}", "status": 200} for index in range(count)
        ],
        "site": {
            "tech_detect": {
                "technologies": [
                    {"category": "server", "name": "Fixture", "evidence": "Saved header"}
                ]
            }
        },
        "suppressed_issues": [
            {
                "id": "excluded",
                "check": "TITLE_MISSING",
                "suppression": {"rule_id": "rule", "reason": "Intentional"},
            }
        ],
        "summary": {
            "pages_checked": count,
            "findings_total": count,
            "findings_by_severity": {"warning": count},
            "severity_note": "Saved coverage basis",
            "checks_disabled": [{"id": "H1_MISSING", "reason": "Operator choice"}],
            "tools_failed": [{"tool": "TITLE_LONG", "error": "Input unavailable"}],
            "finding_exclusion_policy": [{"id": "rule", "reason": "Intentional"}],
            "finding_exclusions": {"suppressed_total": 1},
            "project_coverage": {
                "project": {"site": "example.test", "uuid": "fixture"},
                "status": {
                    "state": "partial",
                    "revision": 1,
                    "items": [{"id": "check", "title": "Saved check", "state": "not_run"}],
                },
            },
            "evidence_contract": {
                "scan_identity_state": "unavailable",
                "scan_identity_reason": "Not retained",
            },
        },
    }


class _Stream:
    def __init__(self, document):
        self.document = document

    def get(self, key, default=None):
        value = self.document.get(key, default)
        return iter(value) if key in {"findings", "pages"} else value


@pytest.mark.parametrize("count", [0, 3])
def test_all_sheet_dimensions_match_emitted_cells_without_read_only_recalculation(
    tmp_path, frozen_office_clock, count
):
    from openpyxl import load_workbook

    document = _document(count)
    materialized = tmp_path / "materialized.xlsx"
    streamed = tmp_path / "streamed.xlsx"
    xlsx.write(document, materialized)
    xlsx.write_stream(_Stream(document), streamed)
    assert materialized.read_bytes() == streamed.read_bytes()
    regular = load_workbook(streamed)
    read_only = load_workbook(streamed, read_only=True)
    try:
        assert set(read_only.sheetnames) == {
            "Summary",
            "Findings",
            "Pages",
            "Technologies",
            "Finding Exclusions",
            "Suppressed Findings",
            "Project Coverage",
            "Evidence coverage",
        }
        for sheet in regular:
            actual = read_only[sheet.title]
            assert actual.max_row == sheet.max_row, sheet.title
            assert actual.max_column == sheet.max_column, sheet.title
            assert actual.calculate_dimension() == sheet.calculate_dimension(), sheet.title
            assert list(actual.values) == list(sheet.values), sheet.title
        assert read_only["Findings"].max_row == read_only["Pages"].max_row == count + 1
        assert len(regular["Summary"]._charts) == 1
        assert read_only["Findings"].max_column == 10
        assert read_only["Pages"].max_column == 11
    finally:
        regular.close()
        read_only.close()


def test_failed_archive_open_closes_xml_generators_and_removes_tempfiles(tmp_path, monkeypatch):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    baseline = set(ALL_TEMP_FILES)
    ignored = []
    monkeypatch.setattr(sys, "unraisablehook", ignored.append)
    with pytest.raises(OSError):
        xlsx.write(_document(1), tmp_path)
    gc.collect()
    assert not ignored
    assert set(ALL_TEMP_FILES) == baseline


@pytest.mark.parametrize("count", [0, 1, 2, 3])
def test_split_sheets_have_exact_rows_dimensions_filters_and_index(tmp_path, monkeypatch, count):
    from openpyxl import load_workbook

    monkeypatch.setattr(xlsx, "MAX_DATA_ROWS", 2)
    document = _document(count)
    for item in document["findings"]:
        item["client_locations"] = [item["url"] + "#exact-evidence"]
    target = tmp_path / "split.xlsx"
    result = xlsx.write_stream(_Stream(document), target)
    index = json.loads(target.with_suffix(".xlsx.index.json").read_text())
    assert result["index"] == str(target.with_suffix(".xlsx.index.json"))
    assert index["state"] == "complete"
    assert index["workbook_sha256"] == hashlib.sha256(target.read_bytes()).hexdigest()
    assert index["source_population"] == {"findings": count, "pages": count}
    workbook = load_workbook(target)
    reader = load_workbook(target, read_only=True)
    try:
        for name, column_count in (("Findings", 10), ("Pages", 11)):
            parts = [part for part in index["ranges"] if part["table"] == name]
            assert len(parts) == max(1, (count + 1) // 2)
            assert sum(part["rows"] for part in parts) == count
            emitted = []
            for number, part in enumerate(parts, 1):
                title = name if number == 1 else f"{name} {number}"
                assert part["worksheet"] == title
                assert reader[title].max_row == part["rows"] + 1
                assert reader[title].max_column == column_count
                assert workbook[title].freeze_panes == "A2"
                rows = list(reader[title].values)[1:]
                emitted.extend(rows)
                assert len(rows) == part["rows"]
                if rows:
                    assert part["first_record"] == (number - 1) * 2 + 1
                    assert part["last_record"] == min(number * 2, count)
                    assert workbook[title].auto_filter.ref.endswith(str(len(rows) + 1))
            url_column = 1 if name == "Findings" else 0
            assert [row[url_column] for row in emitted] == [
                f"https://example.test/{n}" for n in range(count)
            ]
            if name == "Findings":
                assert [row[8] for row in emitted] == [
                    f"https://example.test/{n}#exact-evidence" for n in range(count)
                ]
        assert "Evidence coverage" in index["coverage_worksheets"]
        assert len(workbook["Summary"]._charts) == 1
    finally:
        workbook.close()
        reader.close()


@pytest.mark.parametrize("location", ["finding", "page", "note", "url", "failure", "technology"])
@pytest.mark.parametrize("value", ["x" * 32768, "=" + "x" * 32766])
def test_excel_cell_overflow_is_rejected_before_openpyxl_can_clip(tmp_path, location, value):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    document = _document(1)
    container, key = {
        "finding": (document["findings"][0], "client_observation"),
        "page": (document["pages"][0], "title"),
        "note": (document["summary"], "severity_note"),
        "url": (document, "url"),
        "failure": (document["summary"]["tools_failed"][0], "error"),
        "technology": (document["site"]["tech_detect"]["technologies"][0], "name"),
    }[location]
    container[key] = value
    target = tmp_path / "too-wide.xlsx"
    baseline = set(ALL_TEMP_FILES)
    with pytest.raises(ValueError, match=r"32767-character limit.*CSV or JSON"):
        xlsx.write(document, target)
    assert not list(tmp_path.iterdir())
    assert set(ALL_TEMP_FILES) == baseline


@pytest.mark.parametrize("value", ["x" * 32767, "=" + "x" * 32765])
def test_exact_cell_limit_roundtrips_after_formula_defense(tmp_path, value):
    from openpyxl import load_workbook

    document = _document(1)
    document["pages"][0]["title"] = value
    target = tmp_path / "boundary.xlsx"
    xlsx.write(document, target)
    with target.open("rb") as source:
        workbook = load_workbook(source, read_only=True)
        cell = workbook["Pages"]["C2"]
        assert cell.value == ("'" + value if value.startswith("=") else value)
        assert len(cell.value) == 32767
        assert cell.data_type == "s"
        workbook.close()


@pytest.mark.parametrize("stage", ["archive", "index"])
def test_failed_publication_preserves_previous_workbook_and_index(tmp_path, monkeypatch, stage):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    target = tmp_path / "existing.xlsx"
    sidecar = target.with_suffix(".xlsx.index.json")
    target.write_bytes(b"previous workbook")
    sidecar.write_bytes(b"previous index")
    if stage == "archive":
        from zipfile import ZipFile

        original = ZipFile.writestr
        calls = 0

        def fail(self, *args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("injected ENOSPC")
            return original(self, *args, **kwargs)

        monkeypatch.setattr(ZipFile, "writestr", fail)
    else:
        original = os.replace

        def fail(source, destination):
            if destination == sidecar:
                raise OSError("injected index publication failure")
            return original(source, destination)

        monkeypatch.setattr(os, "replace", fail)
    baseline = set(ALL_TEMP_FILES)
    with pytest.raises(OSError, match="injected"):
        xlsx.write(_document(1), target)
    assert target.read_bytes() == b"previous workbook"
    assert sidecar.read_bytes() == b"previous index"
    assert sorted(p.name for p in tmp_path.iterdir()) == [
        "existing.xlsx",
        "existing.xlsx.index.json",
    ]
    assert set(ALL_TEMP_FILES) == baseline


@pytest.mark.parametrize("after_start", [False, True])
def test_low_disk_rejects_without_partial_outputs_or_spools(tmp_path, monkeypatch, after_start):
    import shutil
    from types import SimpleNamespace

    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    calls = 0

    def disk_usage(_path):
        nonlocal calls
        calls += 1
        return SimpleNamespace(free=(1 << 40) if after_start and calls == 1 else 0)

    monkeypatch.setattr(shutil, "disk_usage", disk_usage)
    baseline = set(ALL_TEMP_FILES)
    with pytest.raises(OSError, match="insufficient free disk"):
        xlsx.write(_document(1), tmp_path / "low-disk.xlsx")
    assert not list(tmp_path.iterdir())
    assert set(ALL_TEMP_FILES) == baseline


def test_one_shot_rows_fail_conservation_instead_of_emitting_header_only(tmp_path):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    document = _document(3)
    document["findings"] = iter(document["findings"])
    baseline = set(ALL_TEMP_FILES)
    with pytest.raises(ValueError, match="row conservation failed"):
        xlsx.write(document, tmp_path / "one-shot.xlsx")
    assert not list(tmp_path.iterdir())
    assert set(ALL_TEMP_FILES) == baseline


def test_row_iteration_interruption_removes_every_spool(tmp_path):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    def interrupted():
        yield {"url": "https://example.test/one"}
        raise KeyboardInterrupt

    document = _document(1)
    document["pages"] = interrupted()
    baseline = set(ALL_TEMP_FILES)
    with pytest.raises(KeyboardInterrupt):
        xlsx.write(document, tmp_path / "interrupted.xlsx")
    assert not list(tmp_path.iterdir())
    assert set(ALL_TEMP_FILES) == baseline
