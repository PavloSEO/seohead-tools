#!/usr/bin/env python3
"""Exercise one complete synthetic native-crawl route at production-scale counts.

The fixture is an owned ``.test`` origin behind ``httpx.MockTransport``.  It
uses the ordinary native handler, robots fetch, sitemap parser and loader,
SQLite collector, analyzer, audit.v2 writer and read-only consumers.  It never
opens a socket and it never inserts rows directly into a scan database.

This is deliberately a manual capacity gate, not a normal test-suite fixture:
the 50k, 100k and 1m stages write real retained artifacts and can take material
time and disk.  Results belong under ``Work/tools/reports/`` and a non-zero
exit means a requested stage did not prove readiness.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import itertools
import json
import math
import os
import platform
import random
import sqlite3
import subprocess
import sys
import time
from collections import Counter
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from unittest.mock import patch

import httpx

# Running ``python scripts/...`` otherwise puts only ``scripts`` ahead of the
# shared editable environment, which can import ``seohead`` from another
# checkout. Every capacity artifact must be bound to this worktree.
PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

HOST = "million-crawl.test"
START_URL = f"https://{HOST}/p/0"
SITEMAP_INDEX = f"https://{HOST}/sitemap-index.xml"
RUNTIME_KEYS = ("python", "sqlite", "httpx", "lxml", "beautifulsoup4")


class SyntheticOrigin:
    """A deterministic HTML, robots and sitemap origin without socket I/O."""

    def __init__(
        self,
        pages: int,
        shard_size: int,
        interrupt_after: int | None = None,
        *,
        links_per_page: int = 1,
        forms_per_page: int = 0,
        body_padding_bytes: int = 0,
        body_profile: str = "padding",
        base_url: str = f"https://{HOST}",
        h1_families: int = 101,
    ) -> None:
        if type(h1_families) is not int or h1_families < 1:
            raise ValueError("h1_families must be a positive integer")
        self.h1_families = h1_families
        if pages < 1:
            raise ValueError("pages must be positive")
        if shard_size < 1:
            raise ValueError("shard_size must be positive")
        if (
            not 1 <= links_per_page <= 150
            or not 0 <= forms_per_page <= 128
            or not 0 <= body_padding_bytes <= 65536
        ):
            raise ValueError("fixture density is outside its declared bounds")
        if body_profile not in {"padding", "catalogue-v1"}:
            raise ValueError("unknown synthetic body profile")
        self.base_url = base_url.rstrip("/")
        self.body_profile = body_profile
        self.links_per_page = min(links_per_page, pages)
        self.forms_per_page = forms_per_page
        self.body_padding_bytes = body_padding_bytes
        self.pages = pages
        self.shard_size = shard_size
        self.interrupt_after = interrupt_after
        self.page_requests = 0
        self.requests: Counter[str] = Counter()

    @property
    def shard_count(self) -> int:
        return (self.pages + self.shard_size - 1) // self.shard_size

    def resume(self) -> None:
        self.interrupt_after = None

    @staticmethod
    def _xml(tag: str, values: Iterator[str]) -> bytes:
        return (
            "<?xml version='1.0' encoding='UTF-8'?>"
            + f"<{tag} xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'>"
            + "".join(values)
            + f"</{tag}>"
        ).encode("utf-8")

    def _index(self) -> bytes:
        return self._xml(
            "sitemapindex",
            (
                f"<sitemap><loc>{self.base_url}/sitemaps/{shard}.xml</loc></sitemap>"
                for shard in range(self.shard_count)
            ),
        )

    def _shard(self, shard: int) -> bytes:
        first = shard * self.shard_size
        last = min(first + self.shard_size, self.pages)
        if shard < 0 or first >= self.pages:
            return b""
        return self._xml(
            "urlset",
            (
                f"<url><loc>{self.base_url}/p/{page}</loc><lastmod>2026-10-04</lastmod>"
                "<changefreq>weekly</changefreq><priority>0.7</priority></url>"
                for page in range(first, last)
            ),
        )

    def _body_extra(self, page: int) -> str:
        if self.body_profile == "padding":
            return "<p>" + ("x" * self.body_padding_bytes) + "</p>"
        rng = random.Random(page)
        materials = (
            "aluminium",
            "steel",
            "polymer",
            "ceramic",
            "composite",
            "glass",
            "copper",
            "silicone",
        )
        finishes = (
            "matte",
            "polished",
            "brushed",
            "coated",
            "textured",
            "anodized",
            "satin",
            "natural",
        )
        roles = (
            "connector",
            "housing",
            "bracket",
            "panel",
            "adapter",
            "frame",
            "terminal",
            "mount",
        )
        parts = []
        size = 0
        while size < self.body_padding_bytes:
            ordinal = len(parts)
            sku = hashlib.sha256(f"synthetic-catalogue:{page}:{ordinal}".encode()).hexdigest()[:24]
            width, height, depth = (rng.randrange(10, 900) for _ in range(3))
            material, finish, role = rng.choice(materials), rng.choice(finishes), rng.choice(roles)
            metadata = json.dumps(
                {
                    "fixture": True,
                    "sku": sku,
                    "batch": rng.randrange(100000, 999999),
                    "dimensions_mm": [width, height, depth],
                    "material": material,
                    "finish": finish,
                },
                separators=(",", ":"),
            )
            section = (
                f'<section data-reference="{sku}"><h2>Specification {ordinal + 1}: {finish} {role}</h2>'
                f"<p>This synthetic catalogue record describes a {material} {role} with a {finish} finish. "
                f"The measured fixture dimensions are {width} by {height} by {depth} millimetres. "
                "Every entry is generated for repeatable storage and parsing measurements; it is not a commercial offer.</p>"
                f"<table><tr><th>Reference</th><td>{sku}</td></tr><tr><th>Material</th><td>{material}</td></tr>"
                f"<tr><th>Envelope</th><td>{width} / {height} / {depth} mm</td></tr></table>"
                f'<script type="application/json">{metadata}</script></section>'
            )
            parts.append(section)
            size += len(section.encode("utf-8"))
        return '<article data-fixture="catalogue-v1">' + "".join(parts) + "</article>"

    def _page(self, page: int) -> bytes:
        if page < 0 or page >= self.pages:
            return b"<html><head><title>Absent</title></head><body>Absent</body></html>"
        links = "".join(
            f"<a href='/p/{(page + hop) % self.pages}'>Next catalogue page {hop}</a>"
            for hop in range(1, self.links_per_page + 1)
        )
        title = "" if page % 997 == 0 else f"Synthetic catalogue page {page}"
        description = "" if page % 991 == 0 else f"Deterministic description for page {page}."
        canonical = f"{self.base_url}/p/{page}"
        extra = "".join(
            f"<form method='post' action='/lead/{form}'><input name='email'></form>"
            for form in range(self.forms_per_page)
        )
        if page % 4096 == 0:
            extra += "<img src='/assets/pixel.png'><form method='post' action='/lead'><input name='email'></form>"
        return (
            "<!doctype html><html lang='en'><head>"
            + (f"<title>{title}</title>" if title else "")
            + (f"<meta name='description' content='{description}'>" if description else "")
            + f"<link rel='canonical' href='{canonical}'>"
            + "<meta name='viewport' content='width=device-width, initial-scale=1'>"
            + "</head><body><main>"
            + f"<h1>Product family {page % self.h1_families}</h1><p>Owned synthetic content for URL {page}.</p>"
            + links
            + extra
            + self._body_extra(page)
            + "</main></body></html>"
        ).encode("utf-8")

    @staticmethod
    def _response(
        request: httpx.Request, status: int, content: bytes = b"", **headers: str
    ) -> httpx.Response:
        """Keep mock bytes unread for the collector's real streaming path."""
        return httpx.Response(
            status,
            stream=httpx.ByteStream(content),
            headers=headers,
            request=request,
        )

    def handle(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if path == "/robots.txt":
            self.requests["robots"] += 1
            content = (
                f"User-agent: *\nAllow: /\nSitemap: {self.base_url}/sitemap-index.xml\n"
            ).encode()
            return self._response(request, 200, content, **{"content-type": "text/plain"})
        if path == "/sitemap-index.xml":
            self.requests["sitemap_index"] += 1
            return self._response(
                request, 200, self._index(), **{"content-type": "application/xml"}
            )
        if path.startswith("/sitemaps/") and path.endswith(".xml"):
            self.requests["sitemap_shard"] += 1
            try:
                shard = int(path.rsplit("/", 1)[-1].removesuffix(".xml"))
            except ValueError:
                return self._response(request, 404)
            content = self._shard(shard)
            return self._response(
                request,
                200 if content else 404,
                content,
                **{"content-type": "application/xml"},
            )
        if path.startswith("/p/"):
            self.page_requests += 1
            if self.interrupt_after is not None and self.page_requests >= self.interrupt_after:
                raise KeyboardInterrupt("intentional synthetic collection interruption")
            try:
                page = int(path.rsplit("/", 1)[-1])
            except ValueError:
                return self._response(request, 404)
            status = 404 if page and page % 100_003 == 0 else 200
            return self._response(
                request,
                status,
                self._page(page),
                **{
                    "content-type": "text/html; charset=utf-8",
                    "cache-control": "max-age=60",
                },
            )
        return self._response(request, 404)


def _body_profile_summary(origin: SyntheticOrigin) -> dict[str, Any]:
    histogram: Counter[int] = Counter()
    digest = hashlib.sha256()
    sizes = []
    for page in range(min(origin.pages, 32)):
        body = origin._page(page)
        histogram.update(body)
        digest.update(len(body).to_bytes(8, "big"))
        digest.update(body)
        sizes.append(len(body))
    total = sum(sizes)
    entropy = -sum((count / total) * math.log2(count / total) for count in histogram.values())
    return {
        "profile": origin.body_profile,
        "sampled_pages": len(sizes),
        "sample_rule": "first up to 32 generated fixture pages, including markup and scripts",
        "sample_sha256": digest.hexdigest(),
        "sample_min_body_bytes": min(sizes),
        "sample_max_body_bytes": max(sizes),
        "sample_mean_body_bytes": round(total / len(sizes), 2),
        "sample_byte_entropy_bits": round(entropy, 6),
        "entropy_note": "Empirical byte-frequency entropy of the declared sample, not a claim of real-site representativeness.",
    }


def _peak_rss_mib() -> float:
    from resource import RUSAGE_SELF, getrusage

    rss = float(getrusage(RUSAGE_SELF).ru_maxrss)
    return round(rss / (1024 * 1024) if sys.platform == "darwin" else rss / 1024, 2)


def _disk_bytes(path: Path) -> dict[str, int]:
    candidates = {
        "scan": path,
        "scan_wal": Path(str(path) + "-wal"),
        "scan_shm": Path(str(path) + "-shm"),
        "audit_v2": path.with_name(path.name + ".audit-v2.sqlite"),
    }
    result = {
        name: candidate.stat().st_size if candidate.exists() else 0
        for name, candidate in candidates.items()
    }
    result["total"] = sum(result.values())
    return result


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _revision() -> str:
    return subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=PROJECT_ROOT, check=True, text=True, capture_output=True
    ).stdout.strip()


def _runtime() -> dict[str, str]:
    import bs4
    import lxml

    return {
        "python": platform.python_version(),
        "sqlite": sqlite3.sqlite_version,
        "httpx": httpx.__version__,
        "lxml": getattr(lxml, "__version__", "unknown"),
        "beautifulsoup4": getattr(bs4, "__version__", "unknown"),
    }


def _loaded_code() -> dict[str, dict[str, str]]:
    """Bind a capacity result to the modules Python actually imported."""
    import hashlib
    import importlib

    modules = (
        "seohead.crawl.settings",
        "seohead.crawl.collect",
        "seohead.tools.parser",
        "seohead.tools.sitemap",
        "seohead.storage.audit_v2",
        "seohead.sf.tasks",
        "seohead.storage.scan_export",
        "seohead.sf.core.inlinks",
        "seohead.servers.scan_handlers",
        "seohead.crawl.sqlite_adapter",
        "seohead.storage.native_scan",
    )
    loaded: dict[str, dict[str, str]] = {}
    for name in modules:
        path = Path(str(importlib.import_module(name).__file__)).resolve()
        if PROJECT_ROOT not in (path, *path.parents):
            raise RuntimeError(f"capacity run imported {name} outside this worktree: {path}")
        loaded[name] = {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    return loaded


def _loaded_callable_code() -> dict[str, dict[str, str]]:
    """Record the callable bodies loaded before the crawl begins, not later files."""
    import hashlib
    import inspect
    import marshal

    from seohead.crawl import sqlite_adapter
    from seohead.servers import handlers, scan_handlers
    from seohead.sf.core import inlinks

    callables = {
        "inlinks._emit_discovery_paths": inlinks._emit_discovery_paths,
        "handlers._audit_crawl_result": handlers._audit_crawl_result,
        "scan_handlers.crawl_site_scan": scan_handlers.crawl_site_scan,
        "sqlite_adapter.crawl_to_scan": sqlite_adapter.crawl_to_scan,
    }
    return {
        name: {
            "file": str(Path(str(callable.__code__.co_filename)).resolve()),
            "source_sha256": hashlib.sha256(inspect.getsource(callable).encode()).hexdigest(),
            "code_sha256": hashlib.sha256(marshal.dumps(callable.__code__)).hexdigest(),
        }
        for name, callable in callables.items()
    }


@contextmanager
def _discovery_path_trace() -> Iterator[list[dict[str, int]]]:
    """Observe the real discovery-path branch without changing emitted findings."""
    from seohead.sf.core import inlinks

    original = inlinks._emit_discovery_paths
    trace: list[dict[str, int]] = []

    def observed(ctx, path_for, *args, **kwargs):
        depths = args[0] if args else kwargs.get("depths")
        population = 0
        resolved = 0
        deep = 0
        if depths is not None:
            for page in ctx.indexable_html_pages():
                population += 1
                key = inlinks.norm_url(page.url)
                depth = depths(key) if callable(depths) else depths.get(key)
                resolved += int(depth is not None)
                deep += int(depth is not None and depth > inlinks._MAX_PATH_URLS)
        path_calls = 0
        longest = 0

        def recorded_path(key):
            nonlocal path_calls, longest
            path = path_for(key)
            path_calls += 1
            longest = max(longest, len(path or ()))
            return path

        original(ctx, recorded_path, *args, **kwargs)
        trace.append(
            {
                "indexable_pages": population,
                "resolved_depths": resolved,
                "missing_depths": population - resolved,
                "deep_depths": deep,
                "path_for_calls": path_calls,
                "longest_path_for_result": longest,
            }
        )

    inlinks._emit_discovery_paths = observed
    try:
        yield trace
    finally:
        inlinks._emit_discovery_paths = original


@contextmanager
def _synthetic_transport(origin: SyntheticOrigin) -> Iterator[None]:
    """Inject one transport beneath real fetch and sitemap code for this run only."""
    from seohead.crawl import collect, sqlite_adapter
    from seohead.sf.core import sitemap_coverage
    from seohead.tools import sitemap

    @contextmanager
    def client_context(_settings: dict[str, Any], _fetcher: Any, _proxy_route: Any = None):
        client = httpx.Client(transport=httpx.MockTransport(origin.handle), follow_redirects=False)
        try:
            yield client
        finally:
            client.close()

    def sitemap_client(*_args: Any, **_kwargs: Any) -> tuple[httpx.Client, bool]:
        return httpx.Client(
            transport=httpx.MockTransport(origin.handle), follow_redirects=True
        ), False

    with (
        patch.object(sqlite_adapter, "_client_context", client_context),
        # ``fetch_one`` normally validates a DNS target before it reaches an
        # opened client. The mock client opens no socket, and this reserved
        # origin has no DNS record to validate; all parsing/dispatch/storage
        # code below this boundary remains the production implementation.
        patch.object(collect, "validate_url", lambda _url, **_kwargs: None),
        patch.object(collect, "pinned_target", lambda url, **_kwargs: (url, {}, {})),
        patch.object(sitemap, "http_client", sitemap_client),
        patch.object(sitemap_coverage, "http_client", sitemap_client),
        # The transport is the owned origin; this is the only no-DNS test seam.
        patch.object(sitemap_coverage, "validate_url", lambda _url: None),
        patch(
            "socket.socket.connect",
            side_effect=AssertionError("synthetic stage attempted socket I/O"),
        ),
    ):
        yield


def _settings(pages: int, *, concurrency: int = 8) -> dict[str, Any]:
    from seohead.crawl.settings import MAX_REQUESTS_CEILING, load

    return load(
        overrides={
            "limits.max_urls": pages,
            # One request per page plus robots, sitemap shards and bounded retries.
            "limits.max_requests": min(MAX_REQUESTS_CEILING, pages + max(10_000, pages // 10)),
            "limits.max_depth": 1,
            "speed.min_delay_seconds": 0,
            "speed.concurrency": concurrency,
            "robots.policy": "respect",
            "sitemaps.auto_discover": False,
            "cache.mode": "off",
            "storage.min_free_bytes": 12 * 1024**3,
        }
    )


def _assert_conservation(scan: Path, pages: int) -> dict[str, int]:
    from seohead.storage import open_scan

    con = open_scan(scan, require_audit=False)
    try:
        counts = {
            name: int(con.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0])
            for name in ("pages", "links", "forms", "frontier")
        }
        counts["sitemap_members"] = int(
            con.execute(
                "SELECT COUNT(*) FROM context_items WHERE kind='sitemap_declared_url'"
            ).fetchone()[0]
        )
        counts["retained_page_documents"] = con.execute(
            "SELECT COUNT(*) FROM pages p JOIN documents d ON d.document_id=p.document_id WHERE d.body_state='complete' AND d.body_sha256 IS NOT NULL"
        ).fetchone()[0]
        for table in ("responses", "documents", "bodies"):
            counts[table] = con.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        sizes = con.execute(
            "SELECT COALESCE(SUM(decoded_bytes),0),COALESCE(SUM(stored_bytes),0),COALESCE(MIN(decoded_bytes),0),COALESCE(MAX(decoded_bytes),0) FROM bodies"
        ).fetchone()
        counts.update(
            zip(
                (
                    "body_decoded_bytes",
                    "body_stored_bytes",
                    "smallest_body_bytes",
                    "largest_body_bytes",
                ),
                sizes,
                strict=True,
            )
        )
        frontier = dict(
            con.execute("SELECT state,COUNT(*) FROM frontier GROUP BY state").fetchall()
        )
    finally:
        con.close()
    if (
        counts["pages"] != pages
        or counts["sitemap_members"] != pages
        or counts["retained_page_documents"] != pages
    ):
        raise AssertionError(f"retained population disagrees with declared scope: {counts}")
    if any(counts[name] != pages for name in ("responses", "documents", "bodies")):
        raise AssertionError(f"response/document/body population disagrees with capture: {counts}")
    if frontier.get("done", 0) != pages or frontier.get("queued", 0) or frontier.get("inflight", 0):
        raise AssertionError(f"frontier is not complete: {frontier}")
    return {**counts, **{f"frontier_{state}": int(count) for state, count in frontier.items()}}


def _collector_summary(result: dict[str, Any]) -> dict[str, Any]:
    """Persist only bounded acceptance facts, never the large audit document."""
    summary = result.get("summary") if isinstance(result.get("summary"), dict) else {}
    return {
        key: result.get(key)
        for key in (
            "scan",
            "urls_collected",
            "links_collected",
            "forms_collected",
            "partial",
            "finish_reason",
            "resumed",
            "audit_available",
            "audit_reason",
            "finalized",
            "corpus_partial",
        )
    } | {
        "audit_totals": summary.get("totals"),
        "audit_coverage": summary.get("check_coverage"),
        "audit_health_score": summary.get("health_score"),
    }


def _recheck_consumers(scan: Path, output: Path, revision: str, pages: int) -> dict[str, Any]:
    """Prove refusal for reused evidence and resolution from a fresh one-page capture."""
    from datetime import datetime, timezone

    from seohead.servers import handlers
    from seohead.servers.scan_handlers import crawl_site_scan
    from seohead.storage.audit_v2 import AuditV2Reader

    with AuditV2Reader(scan) as reader:
        finding_id = next(
            item["id"]
            for item in reader.iter_collection("/issues")
            if item["check"] == "TITLE_MISSING" and item["target_url"] == START_URL
        )
        generated = reader.header["run"]["generated_at"]
    negative = handlers.verify_fixes(
        baseline=str(scan),
        finding_ids=[finding_id],
        after=str(scan),
        out_dir=str(output / "same-observation"),
    )
    if [item["status"] for item in negative["findings"]] != ["not_verifiable"]:
        raise AssertionError("same-observation recheck falsely verified a fix")

    class FixedOrigin(SyntheticOrigin):
        def _page(self, page):
            return (
                super()
                ._page(page)
                .replace(b"</head>", b"<title>Corrected synthetic page title</title></head>")
            )

    baseline_time = datetime.fromisoformat(generated.replace("Z", "+00:00")).timestamp()
    delay = baseline_time + 1.01 - datetime.now(timezone.utc).timestamp()
    if delay > 0:
        time.sleep(delay)
    settings = _settings(pages)
    settings["limits"]["max_urls"] = 1
    fresh_path = output / "fresh-observation.sqlite"
    with _synthetic_transport(FixedOrigin(1, 1)):
        fresh = crawl_site_scan(
            START_URL, scan_out=str(fresh_path), settings=settings, producer_build=revision
        )
    if not fresh.get("audit_available") or not fresh.get("finalized") or fresh.get("partial"):
        raise AssertionError("fresh bounded recheck observation did not complete")
    positive = handlers.verify_fixes(
        baseline=str(scan),
        finding_ids=[finding_id],
        after=str(fresh_path),
        out_dir=str(output / "fresh-recheck"),
    )
    if [item["status"] for item in positive["findings"]] != ["resolved"]:
        raise AssertionError(
            f"fresh local title correction did not resolve selected finding: {positive['findings']!r}"
        )
    return {
        "same_observation": negative["summary"],
        "fresh_observation": positive["summary"],
        "negative_path": negative["verification"],
        "positive_path": positive["verification"],
    }


@contextmanager
def _consumer_phase(output: Path, phases: dict, name: str):
    started = time.monotonic()
    phases[name] = {"state": "running", "started_monotonic": started}

    def persist():
        temporary = output / ".progress.tmp"
        temporary.write_text(
            json.dumps({"schema": "seohead.consumer-progress.v1", "phases": phases}, indent=2)
            + "\n",
            encoding="utf-8",
        )
        temporary.replace(output / "progress.json")

    persist()
    try:
        yield
    except BaseException as exc:
        phases[name].update(state="failed", error=f"{type(exc).__name__}: {str(exc)[:2048]}")
        raise
    else:
        phases[name]["state"] = "returned"
    finally:
        phases[name].update(
            elapsed_seconds=round(time.monotonic() - started, 3),
            cumulative_peak_rss_mib=_peak_rss_mib(),
        )
        persist()


def _csv_rows(path: Path, *, delimiter: str = ",") -> int:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = csv.reader(stream, delimiter=delimiter)
        next(rows)
        return sum(1 for _ in rows)


def _xlsx_rows(path: Path, expected: dict[str, int]) -> dict[str, int]:
    """Read every exported worksheet row and validate its declared partition."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    counts: Counter[str] = Counter()
    try:
        for kind, sheet, first, last, declared in workbook["Partitions"].iter_rows(
            min_row=2, values_only=True
        ):
            actual = sum(1 for _ in workbook[sheet].iter_rows(min_row=2, values_only=True))
            if first != counts[kind] + 1 or last != counts[kind] + actual or actual != declared:
                raise AssertionError(f"XLSX partition does not conserve rows: {kind}, {sheet}")
            counts[kind] += actual
    finally:
        workbook.close()
    if dict(counts) != expected:
        raise AssertionError(f"XLSX readback differs from audit populations: {dict(counts)}")
    return dict(counts)


def _group_evidence(reader) -> dict[str, Any]:
    """Count and hash every ordered member without rebuilding a group array."""
    groups = members = payload = maximum = 0
    digest = hashlib.sha256()
    last_member = None
    for ordinal, group in enumerate(reader.iter_collection("/groups")):
        groups += 1
        group_count = 0
        stream = (
            reader.iter_group_members(ordinal)
            if hasattr(reader, "iter_group_members")
            else group["urls"]
        )
        for member in stream:
            encoded = json.dumps(member, ensure_ascii=False, separators=(",", ":")).encode()
            digest.update(ordinal.to_bytes(8, "big"))
            digest.update(len(encoded).to_bytes(8, "big"))
            digest.update(encoded)
            group_count += 1
            payload += len(encoded)
            maximum = max(maximum, len(encoded))
            last_member = member
        if group_count != len(group["urls"]):
            raise AssertionError("group member stream differs from its declared population")
        if group_count and hasattr(reader, "group_members_page"):
            final = reader.group_members_page(ordinal, offset=group_count - 1, limit=1)
            if final["rows"] != [last_member] or final["has_more"]:
                raise AssertionError("bounded group page cannot recover the final member")
        members += group_count
    return {
        "groups": groups,
        "members": members,
        "ordered_sha256": digest.hexdigest(),
        "member_payload_bytes": payload,
        "largest_member_bytes": maximum,
        "last_member": last_member,
    }


def _consumers(
    scan: Path, output: Path, revision: str, *, comparison_compression: str = "none"
) -> dict[str, Any]:
    from seohead.servers import handlers
    from seohead.sf.tasks import build_tasks_from_audit_v2
    from seohead.storage.audit_v2 import AuditV2Reader
    from seohead.storage.history import snapshot_scan

    output.mkdir(parents=True, exist_ok=True)
    phases = {}
    with _consumer_phase(output, phases, "audit_read"), AuditV2Reader(scan) as reader:
        audit_v2 = {pointer: reader.count(pointer) for pointer in reader.collections}
        if audit_v2.get("/pages") != 0 and audit_v2.get("/issues") is None:
            raise AssertionError("audit.v2 lacks its finding collection")
        payload_bytes = reader.con.execute(
            "SELECT COALESCE(SUM(length(CAST(value_json AS BLOB))),0),COALESCE(MAX(length(CAST(value_json AS BLOB))),0) FROM items"
        ).fetchone()
        header_bytes = reader.con.execute(
            "SELECT length(CAST(header_json AS BLOB)) FROM audit_meta WHERE singleton=1"
        ).fetchone()[0]
        audit_sizes = {
            "collection_payload_bytes": payload_bytes[0],
            "largest_item_bytes": payload_bytes[1],
            "header_bytes": header_bytes,
        }
        member_table = reader.con.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='group_members'"
        ).fetchone()
        member_payload = (
            reader.con.execute(
                "SELECT COALESCE(SUM(length(CAST(value_json AS BLOB))),0) FROM group_members"
            ).fetchone()[0]
            if member_table
            else 0
        )
        audit_sizes["separate_group_member_payload_bytes"] = member_payload
        audit_sizes["total_payload_bytes"] = header_bytes + payload_bytes[0] + member_payload
        summary = reader.header["summary"]
        issue_checks: Counter[str] = Counter()
        occurrences: Counter[str] = Counter()
        for issue in reader.iter_collection("/issues"):
            issue_checks[issue["check"]] += 1
            occurrences[issue["check"]] += issue.get("occurrences_count", 1)
        if not issue_checks or "TITLE_MISSING" not in issue_checks:
            raise AssertionError("fixture did not produce its known nontrivial findings")
        group_evidence = _group_evidence(reader)
        group_members = group_evidence["members"]
        if (
            member_table
            and reader.con.execute("SELECT COUNT(*) FROM group_members").fetchone()[0]
            != group_members
        ):
            raise AssertionError("audit member table and public member stream disagree")
        health = {
            "score": summary.get("health_score"),
            "scope": summary.get("health_score_scope"),
            "check_coverage": summary.get("check_coverage"),
        }
        coverage = health["check_coverage"]
        if not isinstance(coverage, dict) or sum(
            coverage.get(key, 0)
            for key in ("checks_fired", "checks_silent", "checks_skipped", "checks_disabled")
        ) != coverage.get("checks_total"):
            raise AssertionError("retained check coverage denominator is not conserved")
    before_hash = _file_hash(scan)
    companion_path = scan.with_name(scan.name + ".audit-v2.sqlite")
    before_audit_hash = _file_hash(companion_path)
    with _consumer_phase(output, phases, "integrity"):
        integrity = {}
        for name, source in (("scan", scan), ("audit", companion_path)):
            con = sqlite3.connect(source.resolve().as_uri() + "?mode=ro", uri=True)
            try:
                checks = [row[0] for row in con.execute("PRAGMA integrity_check")]
                foreign_keys_valid = con.execute("PRAGMA foreign_key_check").fetchone() is None
                if checks != ["ok"] or not foreign_keys_valid:
                    raise AssertionError(f"{name} failed SQLite integrity: {checks!r}")
                integrity[name] = {"integrity_check": "ok", "foreign_key_check": "ok"}
            finally:
                con.close()
    with _consumer_phase(output, phases, "snapshot"):
        snapshot_path = output / "snapshot.sqlite"
        snapshot_scan(scan, snapshot_path)
        snapshot_counts = _assert_conservation(snapshot_path, audit_v2["/pages"])
        with AuditV2Reader(snapshot_path) as snapshot_reader:
            snapshot_audit = {
                pointer: snapshot_reader.count(pointer) for pointer in snapshot_reader.collections
            }
            if snapshot_audit != audit_v2:
                raise AssertionError("public snapshot did not conserve every audit collection")
    with _consumer_phase(output, phases, "tasks"):
        task_backlog = build_tasks_from_audit_v2(str(scan))
        tasks = task_backlog["tasks"]
        if {task["check"] for task in tasks} != set(issue_checks):
            raise AssertionError("tasks lost a source finding check")
        for task in tasks:
            if task["occurrences"] != occurrences[task["check"]]:
                raise AssertionError("task occurrence total differs from full audit")
            if (
                task["urls"]
                and len(task["urls"]) + task["urls_truncated"] != task["affected_count"]
            ):
                raise AssertionError("task URL coverage hides omitted members")
        task_coverage = {
            "source_findings": sum(issue_checks.values()),
            "source_checks": len(issue_checks),
            "source_occurrences": sum(occurrences.values()),
            "returned_urls": sum(len(task["urls"]) for task in tasks),
            "omitted_urls": sum(task["urls_truncated"] for task in tasks),
        }
    export_populations = {"pages": audit_v2["/pages"], "findings": audit_v2["/issues"]}
    projection = {
        "pages": ["url", "status_code", "title", "canonical"],
        "findings": ["id", "check", "target_url", "severity"],
    }
    with _consumer_phase(output, phases, "export"):
        export = handlers.scan_export(
            input_path=str(scan),
            out=str(output / "scan-export.csv"),
            format="csv",
            records=list(export_populations),
            fields=projection,
        )
        csv_counts = {
            kind: _csv_rows(output / f"scan-export.{kind}.csv") for kind in export_populations
        }
        if csv_counts != export_populations or export.get("counts") != export_populations:
            raise AssertionError("CSV export did not conserve the full source population")
    with _consumer_phase(output, phases, "xlsx_export"):
        xlsx = handlers.scan_export(
            input_path=str(scan),
            out=str(output / "scan-export.xlsx"),
            format="xlsx",
            records=list(export_populations),
            fields=projection,
        )
        if not xlsx.get("ok") or xlsx.get("counts") != export_populations:
            raise AssertionError(f"XLSX export failed: {xlsx!r}")
        xlsx_counts = _xlsx_rows(output / "scan-export.xlsx", export_populations)
    with _consumer_phase(output, phases, "report"):
        report = handlers.report_build(
            audit=str(scan), fmt="csv", out=str(output / "audit-report.csv")
        )
        report_counts = {
            "findings": _csv_rows(output / "audit-report.csv", delimiter=";"),
            "pages": _csv_rows(output / "audit-report.pages.csv", delimiter=";"),
        }
        if report_counts != export_populations:
            raise AssertionError("CSV report did not conserve every source finding and page")
    with _consumer_phase(output, phases, "status"):
        status = handlers.scan_status(input_path=str(scan))
    with _consumer_phase(output, phases, "inspect"):
        inspection = handlers.scan_inspect(
            input_path=str(scan), offset=audit_v2["/pages"] - 1, limit=1
        )
        if (
            len(inspection["rows"]) != 1
            or inspection["rows"][0]["url"] != f"https://{HOST}/p/{audit_v2['/pages'] - 1}"
        ):
            raise AssertionError("bounded inspection did not recover the final retained page")
    with _consumer_phase(output, phases, "diagnosis"):
        diagnosis = handlers.crawl_diagnose(scan=str(scan))
    with _consumer_phase(output, phases, "log_scan"):
        consistency = handlers.log_scan(run=str(scan))
    if not consistency.get("ok") or consistency.get("read", {}).get("pages") != audit_v2["/pages"]:
        raise AssertionError(
            f"consistency consumer failed to read full population: {consistency!r}"
        )
    with _consumer_phase(output, phases, "reanalysis"):
        reanalysis = handlers.scan_reanalyze(
            input_path=str(scan), out=str(output / "reanalysis.sqlite"), producer_build=revision
        )
    compare_options = (
        {"compression": comparison_compression} if comparison_compression != "none" else {}
    )
    with _consumer_phase(output, phases, "compare"):
        comparison = handlers.compare_crawls(
            before=str(scan),
            after=str(output / "reanalysis.sqlite"),
            out_dir=str(output / "comparison"),
            **compare_options,
        )
    from seohead.sf.core.compare_store import iter_compare_rows

    comparison_roundtrip = {
        name: sum(1 for _ in iter_compare_rows(comparison["manifest"], name))
        for name in comparison["files"]
    }
    if any(
        comparison_roundtrip[name] != item["rows"] for name, item in comparison["files"].items()
    ):
        raise AssertionError("comparison package did not round-trip every retained row")
    if (
        comparison.get("conservation", {}).get("state") != "complete"
        or comparison["conservation"]["before_issues"] != audit_v2["/issues"]
    ):
        raise AssertionError(f"bounded comparison did not conserve source findings: {comparison!r}")
    with _consumer_phase(output, phases, "recheck"):
        recheck = _recheck_consumers(scan, output, revision, audit_v2["/pages"])
    after_hash = _file_hash(scan)
    after_audit_hash = _file_hash(companion_path)
    if before_hash != after_hash or before_audit_hash != after_audit_hash:
        raise AssertionError("read-only consumers or reanalysis changed the source scan")
    with AuditV2Reader(output / "reanalysis.sqlite") as derived:
        if derived.count("/pages") != audit_v2["/pages"]:
            raise AssertionError("offline reanalysis did not conserve the page population")
    if not export.get("ok") or not report.get("ok") or not reanalysis.get("audit_available"):
        raise AssertionError(
            f"consumer failure: export={export!r}; report={report!r}; reanalysis={reanalysis!r}"
        )
    return {
        "audit_v2": audit_v2,
        "audit_sizes": audit_sizes,
        "source_sha256_before": before_hash,
        "source_sha256_after": after_hash,
        "source_audit_sha256_before": before_audit_hash,
        "source_audit_sha256_after": after_audit_hash,
        "integrity": integrity,
        "snapshot": {
            "path": str(snapshot_path),
            "conservation": snapshot_counts,
            "audit_v2": snapshot_audit,
        },
        "tasks": task_backlog["summary"],
        "task_coverage": task_coverage,
        "group_members": group_members,
        "group_evidence": group_evidence,
        "health": health,
        "comparison": comparison,
        "comparison_roundtrip": comparison_roundtrip,
        "export": export,
        "csv_readback": csv_counts,
        "xlsx_export": xlsx,
        "xlsx_readback": xlsx_counts,
        "report": report,
        "report_readback": report_counts,
        "status": status,
        "inspection": {
            "offset": inspection["offset"],
            "rows": len(inspection["rows"]),
            "last_url": inspection["rows"][0]["url"],
            "has_more": inspection["has_more"],
        },
        "diagnosis_codes": [item["code"] for item in diagnosis.get("diagnoses", [])],
        "consistency": consistency,
        "reanalysis": reanalysis,
        "recheck": recheck,
    }


def run_stage(
    output: Path,
    *,
    pages: int,
    shard_size: int,
    interrupt_after: int | None,
    consumers: bool,
    links_per_page: int = 1,
    forms_per_page: int = 0,
    body_padding_bytes: int = 0,
    comparison_compression: str = "none",
    body_profile: str = "padding",
    h1_families: int = 101,
) -> dict[str, Any]:
    """Run one measured stage; exceptions intentionally make its status failed."""
    from seohead.servers.scan_handlers import crawl_site_scan
    from seohead.storage import open_scan
    from seohead.storage.native_scan import NativeScan
    from seohead.tools import sitemap

    if pages > sitemap.MAX_URLS:
        raise RuntimeError(
            f"standard sitemap parser limit is {sitemap.MAX_URLS:,}, below requested {pages:,}; "
            "a native 1m crawl through the normal sitemap route is not admitted"
        )
    output.mkdir(parents=True, exist_ok=True)
    scan = output / "native.sqlite"
    origin = SyntheticOrigin(
        pages,
        shard_size,
        interrupt_after,
        links_per_page=links_per_page,
        forms_per_page=forms_per_page,
        body_padding_bytes=body_padding_bytes,
        body_profile=body_profile,
        h1_families=h1_families,
    )
    # A synthetic KeyboardInterrupt must reach the real collector coordinator,
    # not be retained as one failed member of a concurrent fixture batch.
    # The resumed run uses this same settings object, preserving its contract.
    settings = _settings(pages, concurrency=1 if interrupt_after is not None else 8)
    revision = _revision()
    loaded_code = _loaded_code()
    loaded_callable_code = _loaded_callable_code()
    started = time.monotonic()
    checkpoint: dict[str, Any] | None = None
    with _synthetic_transport(origin), _discovery_path_trace() as discovery_path_trace:
        if interrupt_after is not None:
            try:
                interrupted = crawl_site_scan(
                    START_URL,
                    scan_out=str(scan),
                    settings=settings,
                    sitemap=SITEMAP_INDEX,
                    producer_build=revision,
                )
            except KeyboardInterrupt:
                interrupted = {"partial": True, "finish_reason": "interrupted"}
            else:
                if (
                    not interrupted.get("partial")
                    or interrupted.get("finish_reason") != "interrupted"
                ):
                    raise AssertionError("synthetic interruption did not interrupt the collector")
            con = open_scan(scan, require_audit=False)
            try:
                row = con.execute("SELECT lifecycle,finish_reason FROM scan").fetchone()
                committed = int(con.execute("SELECT COUNT(*) FROM pages").fetchone()[0])
                frontier = dict(con.execute("SELECT state,COUNT(*) FROM frontier GROUP BY state"))
            finally:
                con.close()
            if not 0 < committed < pages or row["lifecycle"] == "finished":
                raise AssertionError(
                    f"interruption did not leave a resumable prefix: {committed}, {dict(row)}"
                ) from None
            checkpoint = {
                "pages": committed,
                "lifecycle": row["lifecycle"],
                "finish_reason": row["finish_reason"],
                "frontier": frontier,
            }
            origin.resume()
        result = crawl_site_scan(
            START_URL,
            scan_out=str(scan),
            settings=settings,
            sitemap=SITEMAP_INDEX,
            producer_build=revision,
        )
    finalization_retried = False
    if result.get("audit_available") and not result.get("finalized"):
        # A real interruption can leave the final WAL checkpoint temporarily
        # blocked even after the resumed collector and audit both completed.
        # Re-open through the public native writer and prove that its explicit
        # recovery path finalizes the exact same artifact before calling it done.
        with NativeScan.open(scan) as writer:
            result["finalized"] = writer.finish_capture(reason="finished")
        finalization_retried = True
    if not result.get("audit_available") or not result.get("finalized"):
        raise AssertionError(f"collection did not publish an audit: {result!r}")
    if result.get("partial") or result.get("corpus_partial"):
        raise AssertionError(
            f"requested scope or retained corpus is incomplete: {_collector_summary(result)!r}"
        )
    if interrupt_after is not None and not result.get("resumed"):
        raise AssertionError("interrupted scan was recollected instead of resumed")
    expected_requests = pages + int(interrupt_after is not None)
    if origin.page_requests != expected_requests:
        raise AssertionError(
            f"resume refetched pages: {origin.page_requests} != {expected_requests}"
        )
    counts = _assert_conservation(scan, pages)
    if (
        counts["links"] != pages * origin.links_per_page
        or counts["forms"] != pages * forms_per_page + (pages + 4095) // 4096
    ):
        raise AssertionError(f"retained links/forms disagree with generated source: {counts}")
    record: dict[str, Any] = {
        "status": "passed",
        "acceptance_scope": "producer_and_consumers" if consumers else "producer_only",
        "pages_requested": pages,
        "shard_size": shard_size,
        "body_profile_sample": _body_profile_summary(origin),
        "fixture": {
            "schema": "seohead.synthetic-crawl.v3",
            "body_profile": body_profile,
            "h1_families": h1_families,
            "links_per_page": origin.links_per_page,
            "forms_per_page": forms_per_page,
            "body_padding_bytes": body_padding_bytes,
        },
        "source_revision": revision,
        "runtime": _runtime(),
        "loaded_code": loaded_code,
        "loaded_callable_code": loaded_callable_code,
        # An intentional interrupted prefix is evidence about resume, while
        # this final metric names the finished capture only. Keep both rather
        # than silently merging populations from different lifecycle states.
        "discovery_path_trace": discovery_path_trace[-1:],
        "interruption_discovery_path_trace": discovery_path_trace[:-1],
        "elapsed_seconds": round(time.monotonic() - started, 3),
        "peak_rss_mib": _peak_rss_mib(),
        "disk_bytes": _disk_bytes(scan),
        "transport_requests": dict(origin.requests),
        "page_requests": origin.page_requests,
        "checkpoint": checkpoint,
        "finalization_retried": finalization_retried,
        "collector": _collector_summary(result),
        "conservation": counts,
        "body_storage": {
            "stored_to_decoded_ratio": round(
                counts["body_stored_bytes"] / counts["body_decoded_bytes"], 6
            )
            if counts["body_decoded_bytes"]
            else None,
            "note": "Body-codec ratio only; excludes SQLite indexes, audit and consumer files.",
        },
    }
    record["capture_audit_seconds"] = record["elapsed_seconds"]
    (output / "producer-result.json").write_text(
        json.dumps(
            {
                **record,
                "status": "producer_passed_consumers_pending"
                if consumers
                else "producer_passed_consumers_skipped",
            },
            indent=2,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )
    if consumers:
        record["consumers"] = _consumers(
            scan, output / "consumers", revision, comparison_compression=comparison_compression
        )
    record["elapsed_seconds"] = round(time.monotonic() - started, 3)
    record["peak_rss_mib"] = _peak_rss_mib()
    (output / "result.json").write_text(
        json.dumps(record, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8"
    )
    return record


def run_consumers_only(
    scan: Path, output: Path, *, comparison_compression: str = "none"
) -> dict[str, Any]:
    """Recheck consumers against a preserved synthetic capture without recollection."""
    from contextlib import closing

    from seohead.storage import open_scan

    started = time.monotonic()
    loaded_code, loaded_callables = _loaded_code(), _loaded_callable_code()
    with closing(open_scan(scan, require_audit=False)) as con:
        source = dict(
            con.execute("SELECT writer_revision,start_url FROM scan WHERE singleton=1").fetchone()
        )
    if source["start_url"] != START_URL:
        raise ValueError("capacity consumer retry requires the owned synthetic fixture origin")
    if output.exists() and any(output.iterdir()):
        raise ValueError("consumer retry output must be empty or new")
    output.mkdir(parents=True, exist_ok=True)
    consumers = _consumers(
        scan, output / "consumers", _revision(), comparison_compression=comparison_compression
    )
    result = {
        "status": "passed",
        "mode": "consumers_only",
        "capture_source_revision": source["writer_revision"],
        "consumer_source_revision": _revision(),
        "scan": str(scan),
        "loaded_code": loaded_code,
        "loaded_callable_code": loaded_callables,
        "consumers": consumers,
        "elapsed_seconds": round(time.monotonic() - started, 3),
        "peak_rss_mib": _peak_rss_mib(),
    }
    (output / "result.json").write_text(
        json.dumps(result, indent=2, default=str) + "\n", encoding="utf-8"
    )
    return result


def run_loopback(output: Path, *, pages: int = 8) -> dict[str, Any]:
    """Exercise the owned HTTP origin with the actual guarded client and pacing."""
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    from seohead.crawl.settings import effective_request_rate, load
    from seohead.recon.net import validate_url
    from seohead.servers.scan_handlers import crawl_site_scan

    if not 2 <= pages <= 100:
        raise ValueError("loopback acceptance is a separate 2..100-page transport smoke")
    if output.exists() and any(output.iterdir()):
        raise ValueError("loopback output must be empty or new")
    output.mkdir(parents=True, exist_ok=True)
    received: list[float] = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            received.append(time.monotonic())
            response = origin.handle(httpx.Request("GET", origin.base_url + self.path))
            body = response.read()
            self.send_response(response.status_code)
            for name, value in response.headers.items():
                self.send_header(name, value)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    base_url = f"http://127.0.0.1:{server.server_port}"
    origin = SyntheticOrigin(pages, max(1, pages // 2), base_url=base_url, links_per_page=3)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    started = time.monotonic()
    settings = load(overrides={"limits.max_urls": pages, "limits.max_requests": pages + 20})
    scan = output / "loopback.sqlite"
    try:
        with patch.dict(
            os.environ, {"SEOHEAD_ALLOW_PRIVATE_NETWORKS": "0", "SEOHEAD_ALLOW_PRIVATE_HOSTS": ""}
        ):
            try:
                validate_url(base_url)
            except ValueError:
                private_refused = True
            else:
                raise AssertionError("guard admitted loopback without the owned-host allowance")
        with patch.dict(
            os.environ,
            {"SEOHEAD_ALLOW_PRIVATE_NETWORKS": "0", "SEOHEAD_ALLOW_PRIVATE_HOSTS": "127.0.0.1"},
        ):
            result = crawl_site_scan(
                base_url + "/p/0",
                scan_out=str(scan),
                settings=settings,
                sitemap=base_url + "/sitemap-index.xml",
                producer_build=_revision(),
            )
        if (
            result.get("partial")
            or not result.get("audit_available")
            or not result.get("finalized")
        ):
            raise AssertionError(
                f"guarded HTTP collection did not finish: {_collector_summary(result)}"
            )
        counts = _assert_conservation(scan, pages)
        if (
            counts["links"] != pages * origin.links_per_page
            or counts["forms"] != (pages + 4095) // 4096
        ):
            raise AssertionError("owned HTTP fixture did not retain every parsed link and form")
        from seohead.storage.audit_v2 import AuditV2Reader

        with AuditV2Reader(scan) as reader:
            title_missing = sum(
                item["check"] == "TITLE_MISSING" for item in reader.iter_collection("/issues")
            )
        if title_missing != 1:
            raise AssertionError("owned HTTP fixture lost its known title finding")
        # Dispatch slots are reserved at least ``min_delay`` apart, but each
        # request's arrival timestamp also carries connect and scheduling
        # jitter, so on a loaded host one server-observed pair can dip slightly
        # under the delay floor. Pacing evidence is the aggregate window plus a
        # burst floor that a collapsed throttle cannot meet.
        intervals = [after - before for before, after in itertools.pairwise(received)]
        paced_window = received[-1] - received[0]
        mean_interval = paced_window / len(intervals)
        if (
            effective_request_rate(settings) > 2
            or min(intervals) < 0.2
            or mean_interval < 0.45
        ):
            raise AssertionError(
                "owned HTTP fixture did not preserve the default 2-request/s pacing"
            )
        receipt = {
            "status": "passed",
            "mode": "guarded_loopback",
            "source_revision": _revision(),
            "loaded_code": _loaded_code(),
            "runtime": _runtime(),
            "private_refused_before_allowance": private_refused,
            "allowed_private_host": "127.0.0.1",
            "transport": "real guarded http_client",
            "effective_max_requests_per_second": effective_request_rate(settings),
            "minimum_observed_request_interval_seconds": min(intervals),
            "mean_observed_request_interval_seconds": mean_interval,
            "requests": len(received),
            "known_title_missing_findings": title_missing,
            "conservation": counts,
            "elapsed_seconds": time.monotonic() - started,
            "peak_rss_mib": _peak_rss_mib(),
        }
        (output / "result.json").write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
        return receipt
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def _check_peak_rss(result: dict[str, Any], limit: float, path: Path) -> None:
    if result["peak_rss_mib"] > limit:
        result.update(status="blocked", reason="operating_system_peak_rss_budget")
        path.write_text(json.dumps(result, indent=2, default=str) + "\n", encoding="utf-8")
        raise RuntimeError("stage exceeded the operating-system peak RSS budget")


def _parse_stages(value: str) -> list[int]:
    try:
        stages = [int(item) for item in value.split(",") if item]
    except ValueError as exc:
        raise argparse.ArgumentTypeError("stages must be comma-separated integers") from exc
    if not stages or any(stage < 1 for stage in stages):
        raise argparse.ArgumentTypeError("stages must contain positive counts")
    return stages


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        required=True,
        type=Path,
        help="empty or new report directory outside the repository",
    )
    parser.add_argument("--stages", type=_parse_stages, default=[50_000, 100_000, 1_000_000])
    parser.add_argument("--shard-size", type=int, default=50_000)
    parser.add_argument("--interrupt-after", type=int, default=25_000)
    parser.add_argument(
        "--input-scan",
        type=Path,
        help="retry all consumers on an existing owned synthetic capture; never recollect it",
    )
    parser.add_argument("--skip-consumers", action="store_true")
    parser.add_argument(
        "--loopback-only", action="store_true", help="explicit owned HTTP smoke with default pacing"
    )
    parser.add_argument("--links-per-page", type=int, default=3)
    parser.add_argument("--forms-per-page", type=int, default=1)
    parser.add_argument("--h1-families", type=int, default=1)
    parser.add_argument("--body-padding-bytes", type=int, default=2048)
    parser.add_argument("--comparison-compression", choices=("none", "gzip"), default="none")
    parser.add_argument(
        "--body-profile", choices=("padding", "catalogue-v1"), default="catalogue-v1"
    )
    parser.add_argument("--max-seconds", type=float, default=1800)
    parser.add_argument("--max-rss-mib", type=float, default=4096)
    parser.add_argument("--max-disk-mib", type=float, default=32768)
    parser.add_argument("--min-free-mib", type=float, default=12288)
    parser.add_argument("--stage-worker", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if args.input_scan is not None and args.skip_consumers:
        parser.error("--input-scan cannot skip consumers")
    if args.input_scan is not None and args.loopback_only:
        parser.error("--input-scan and --loopback-only are mutually exclusive")
    if args.interrupt_after < 1:
        parser.error("--interrupt-after must be positive")
    if args.out.exists() and any(args.out.iterdir()):
        parser.error("--out must be an empty or new directory so evidence cannot be mixed")
    status = subprocess.run(
        ["git", "status", "--porcelain"],
        cwd=PROJECT_ROOT,
        check=True,
        text=True,
        capture_output=True,
    )
    if status.stdout.strip():
        parser.error("capacity acceptance requires a clean, frozen source checkout")
    if not args.stage_worker:
        from scripts.capacity_watchdog import supervise

        args.out.mkdir(parents=True, exist_ok=True)
        stages = [None] if args.input_scan is not None or args.loopback_only else args.stages
        results = []
        for pages in stages:
            stage = args.out / (
                str(pages) if pages is not None else "consumers" if args.input_scan else "loopback"
            )
            temporary = stage / "tmp"
            temporary.mkdir(parents=True)
            child_out = stage / "artifacts"
            command = [
                sys.executable,
                str(Path(__file__).resolve()),
                "--stage-worker",
                "--out",
                str(child_out),
            ]
            if args.input_scan:
                command += ["--input-scan", str(args.input_scan.resolve())]
            elif args.loopback_only:
                command.append("--loopback-only")
            else:
                command += [
                    "--stages",
                    str(pages),
                    "--shard-size",
                    str(args.shard_size),
                    "--interrupt-after",
                    str(args.interrupt_after),
                    "--links-per-page",
                    str(args.links_per_page),
                    "--forms-per-page",
                    str(args.forms_per_page),
                    "--h1-families",
                    str(args.h1_families),
                    "--body-padding-bytes",
                    str(args.body_padding_bytes),
                    "--body-profile",
                    args.body_profile,
                ]
                if args.skip_consumers:
                    command.append("--skip-consumers")
            command += [
                "--comparison-compression",
                args.comparison_compression,
                "--max-rss-mib",
                str(args.max_rss_mib),
            ]
            receipt = supervise(
                command,
                cwd=PROJECT_ROOT,
                output=stage / "watchdog",
                disk_dir=stage,
                env={
                    **os.environ,
                    "TMPDIR": str(temporary),
                    "SQLITE_TMPDIR": str(temporary),
                    "TMP": str(temporary),
                    "TEMP": str(temporary),
                },
                max_seconds=args.max_seconds,
                max_rss_mib=args.max_rss_mib,
                max_disk_mib=args.max_disk_mib,
                min_free_mib=args.min_free_mib,
                measured_paths={"temporary_spools": temporary, "artifacts": child_out},
            )
            results.append(
                {
                    "pages_requested": pages,
                    "watchdog": str(stage / "watchdog/watchdog.json"),
                    **receipt,
                }
            )
            if receipt["status"] != "passed":
                break
        manifest = {
            "format": "seohead.million-crawl-acceptance.v2",
            "source_revision": _revision(),
            "results": results,
        }
        (args.out / "manifest.json").write_text(
            json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(manifest, indent=2))
        return (
            0
            if len(results) == len(stages) and all(item["status"] == "passed" for item in results)
            else 1
        )
    if args.loopback_only:
        result = run_loopback(args.out)
        _check_peak_rss(result, args.max_rss_mib, args.out / "result.json")
        return 0
    if args.input_scan is not None:
        if args.skip_consumers:
            parser.error("--input-scan cannot skip consumers")
        try:
            result = run_consumers_only(
                args.input_scan, args.out, comparison_compression=args.comparison_compression
            )
            _check_peak_rss(result, args.max_rss_mib, args.out / "result.json")
        except BaseException:
            import traceback

            args.out.mkdir(parents=True, exist_ok=True)
            (args.out / "consumer-failure.txt").write_text(traceback.format_exc(), encoding="utf-8")
            return 1
        print(json.dumps(result, indent=2, default=str))
        return 0
    args.out.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    for pages in args.stages:
        interruption = min(args.interrupt_after, max(1, pages // 2))
        try:
            results.append(
                run_stage(
                    args.out / str(pages),
                    pages=pages,
                    shard_size=args.shard_size,
                    interrupt_after=interruption,
                    consumers=not args.skip_consumers,
                    links_per_page=args.links_per_page,
                    forms_per_page=args.forms_per_page,
                    body_padding_bytes=args.body_padding_bytes,
                    comparison_compression=args.comparison_compression,
                    body_profile=args.body_profile,
                    h1_families=args.h1_families,
                )
            )
            _check_peak_rss(results[-1], args.max_rss_mib, args.out / str(pages) / "result.json")
        except BaseException as exc:
            import traceback

            (args.out / f"{pages}-failure.txt").write_text(traceback.format_exc(), encoding="utf-8")
            failed = {
                "status": "failed",
                "pages_requested": pages,
                "error": f"{type(exc).__name__}: {exc}",
                "elapsed_seconds": None,
            }
            results.append(failed)
            break
    manifest = {"format": "seohead.million-crawl-acceptance.v1", "results": results}
    (args.out / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return (
        0
        if len(results) == len(args.stages) and all(row["status"] == "passed" for row in results)
        else 1
    )


if __name__ == "__main__":
    raise SystemExit(main())
